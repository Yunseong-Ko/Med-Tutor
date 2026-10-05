#!/usr/bin/env python3
"""부산의대 PMA 공식 출제계획(붙임2·3) → 기계 블루프린트.

학교 체계:
  붙임2  CP번호 × 출제과 × 차수별 문항수 (쿼터 매트릭스)
  붙임3  CP 122개: 평가목표 · 맥락/질환 · 핵심성과 · 구체적성과
차수: 1차=4학년 자체1차(320) · 2차=자체2차(320) · 3차=3학년 자체(280) · 4차=3학년 재시(200) · 5차=4학년 전국(320)

출력: data_private/curriculum/pnu_pma_blueprint.json
  {cps: {cp_no: {name, contexts[], core_outcomes, specific_outcomes}},
   quota: [{cp_no, dept, counts: {1cha..5cha}}], meta}
"""
import argparse
import json
import re
from collections import Counter
from pathlib import Path

import openpyxl

OUT = Path("data_private/curriculum/pnu_pma_blueprint.json")
ROUNDS = ["1cha", "2cha", "3cha", "4cha", "5cha"]
# 붙임2 열: A=CP번호 B=출제과 C=CP명 D=1차A형 E=1차멀티 F=2차A형 G=2차멀티 H=3차 I=4차 J=5차
QUOTA_COLS = {0: "cp_no", 1: "dept", 2: "cp_name", 3: "1cha", 5: "2cha", 7: "3cha", 8: "4cha", 9: "5cha"}


def _count(v) -> int:
    """'예시) 1' 같은 표기·공백·x 를 안전하게 정수로."""
    if v is None:
        return 0
    m = re.search(r"\d+", str(v))
    return int(m.group()) if m else 0


def parse_quota(path: Path):
    """이 시트는 **작성 양식**이다: 과 헤더행에 과별 총량이 있고,
    CP행의 차수별 분배는 각 교실이 채운다(감염내과만 예시로 기입돼 있음).
    → (과별 총량, 과별 CP 메뉴, 기입된 CP 분배)를 모두 보존한다."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["전체"]
    cp_rows, dept_totals, cur_dept = [], {}, ""
    for r in ws.iter_rows(min_row=4, values_only=True):
        a = str(r[0] or "").strip()
        if not a:
            continue
        if not a.isdigit():
            if a.startswith("#"):
                continue
            cur_dept = a                          # 과 헤더행 = 과별 차수 총량
            dept_totals[cur_dept] = {rd: _count(r[c]) for c, rd in
                ((3, "1cha"), (5, "2cha"), (7, "3cha"), (8, "4cha"), (9, "5cha"))}
            continue
        dept = str(r[1] or "").strip() or cur_dept
        counts = {rd: _count(r[c]) for c, rd in
                  ((3, "1cha"), (5, "2cha"), (7, "3cha"), (8, "4cha"), (9, "5cha"))}
        cp_rows.append({"cp_no": int(a), "dept": dept,
                        "cp_name": str(r[2] or "").strip(), "counts": counts})
    return cp_rows, dept_totals


def parse_cps(path: Path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["전체"]
    cps = {}
    for r in ws.iter_rows(min_row=3, values_only=True):
        no = str(r[0] or "").strip()
        if not no.isdigit():
            continue
        contexts = [c.strip() for c in re.split(r"[,，·/]| 및 ", str(r[2] or ""))
                    if c.strip() and len(c.strip()) >= 2]
        cps[int(no)] = {
            "name": str(r[1] or "").strip(),
            "contexts": contexts,
            "core_outcomes": str(r[3] or "").strip(),
            "specific_outcomes": str(r[4] or "").strip(),
        }
    return cps


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--quota", default=str(Path.home() / "Downloads/[붙임2]2026 출제계획표 A형 과별 문항수.xlsx"))
    ap.add_argument("--cps", default=str(Path.home() / "Downloads/[붙임3]CP별 평가목표.xlsx"))
    args = ap.parse_args()

    quota, dept_totals = parse_quota(Path(args.quota))
    cps = parse_cps(Path(args.cps))

    linked = sum(1 for q in quota if q["cp_no"] in cps)
    per_round = Counter()
    for t in dept_totals.values():
        for rd, n in t.items():
            per_round[rd] += n
    depts = sorted({q["dept"] for q in quota if q["dept"]})

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps({
        "schema": "pnu_pma_blueprint.v1",
        "source": "2026 문항출제 공문 붙임2·3",
        "cps": {str(k): v for k, v in sorted(cps.items())},
        "quota": quota,
        "dept_totals": dept_totals,
        "meta": {"rounds": ROUNDS, "departments": depts,
                 "note": "CP행 counts는 교실이 채우는 칸(감염내과만 예시 기입). "
                         "과별 확정 총량은 dept_totals."},
    }, ensure_ascii=False, indent=1), encoding="utf-8")

    print(f"CP 정의 {len(cps)}개 · 쿼터행 {len(quota)}개 (CP 연결 {linked}) · 과 {len(depts)}개")
    print("  차수별 합계:", dict(per_round), "(공문 기준 320/320/280/200/320)")
    miss = [q["cp_no"] for q in quota if q["cp_no"] not in cps]
    if miss:
        print("  ! 붙임3에 없는 CP번호:", sorted(set(miss))[:10])
    print(f"[출력] {OUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
