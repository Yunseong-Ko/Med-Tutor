#!/usr/bin/env python3
"""회수된 종합설문(학생_NN_종합설문.xlsx) 집계 → 논문 데이터 ③.

각 파일에서 9문항 Likert 응답(D5:D13)과 자유의견(종합 의견 행 아래 셀)을 읽어
문항별 평균·표준편차·중앙값·응답분포와 개방형 답변 목록을 산출한다.
검토표 집계(aggregate_review_feedback.py)와 마찬가지로 파일명 번호=응답자.

출력: exports/survey_summary.md + survey_responses.csv (논문 Table 재료)
"""
import argparse
import csv
import statistics
from pathlib import Path

import openpyxl

DEFAULT_DIR = Path("data_private/professor_items/exports/회수_설문")
OUT_DIR = Path("data_private/professor_items/exports")

AREAS = [
    "안면타당도", "내용 포괄성", "임상적 관련성", "임상추론 평가 적절성",
    "난이도 적절성", "문항 완성도", "정답·해설의 교육적 유용성",
    "평가도구 활용 가능성", "학습도구 활용 가능성",
]


def read_form(path: Path):
    ws = openpyxl.load_workbook(path, data_only=True)["종합설문"]
    resp = {}
    for i in range(9):
        v = ws.cell(5 + i, 4).value
        try:
            v = int(v)
        except (TypeError, ValueError):
            v = None
        if v is not None and not 1 <= v <= 5:
            v = None
        resp[i + 1] = v
    open_text = ""
    # "종합 의견" 라벨 행을 찾아 그 아래 셀(자유 기술란)을 읽는다
    for r in range(14, ws.max_row + 1):
        if str(ws.cell(r, 2).value or "") == "종합 의견":
            open_text = str(ws.cell(r + 1, 3).value or "").strip()
            break
    return resp, open_text


def write_outputs(rows, comments):
    """집계 산출물 저장. rows=[{respondent, q1..q9}], comments=[(respondent, text)].

    구글폼 CSV 경로(aggregate_gform_csv.py)와 공유.
    """
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    with (OUT_DIR / "survey_responses.csv").open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=["respondent"] + [f"q{i}" for i in range(1, 10)])
        w.writeheader()
        w.writerows(rows)

    lines = [
        "# 종합 설문 집계 (논문 데이터 ③)", "",
        f"- 응답 {len(rows)}부 · 자유의견 {len(comments)}건", "",
        "| # | 평가영역 | n | 평균 | SD | 중앙값 | 분포(1~5) |",
        "|---|---|---|---|---|---|---|",
    ]
    for i in range(1, 10):
        vals = [r[f"q{i}"] for r in rows if r[f"q{i}"] is not None]
        dist = "·".join(str(sum(1 for v in vals if v == s)) for s in range(1, 6))
        if vals:
            mean = f"{statistics.mean(vals):.2f}"
            sd = f"{statistics.stdev(vals):.2f}" if len(vals) > 1 else "—"
            med = f"{statistics.median(vals):.1f}"
        else:
            mean = sd = med = "—"
        lines.append(f"| {i} | {AREAS[i-1]} | {len(vals)} | {mean} | {sd} | {med} | {dist} |")
    lines += ["", "## 자유의견 (문항 10)", ""]
    for digits, txt in comments:
        lines.append(f"- **응답자 {digits}**: {txt}")
    (OUT_DIR / "survey_summary.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"집계 {len(rows)}부 → {OUT_DIR / 'survey_summary.md'} · survey_responses.csv")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", default=str(DEFAULT_DIR), help="회수된 설문 xlsx 폴더")
    args = ap.parse_args()
    src = Path(args.dir)
    files = sorted(p for p in src.glob("*종합설문*.xlsx") if not p.name.startswith("~$"))
    if not files:
        print(f"설문 파일 없음: {src}")
        return 1

    rows, comments = [], []
    for p in files:
        digits = "".join(ch for ch in p.stem if ch.isdigit()).zfill(2)
        resp, open_text = read_form(p)
        rows.append({"respondent": digits, **{f"q{i}": resp[i] for i in range(1, 10)}})
        if open_text:
            comments.append((digits, open_text))
    write_outputs(rows, comments)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
