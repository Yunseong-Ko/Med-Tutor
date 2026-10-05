#!/usr/bin/env python3
"""
Apply team labeling spreadsheet rows to extracted course-exam JSON files.

This script intentionally avoids printing original question stems or choices.
It matches labels by exam round/source and question number, then writes the
enriched JSON back to data_private for local-only use.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Any


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))


HEADER_MAP = {
    "문항 번호": "question_number",
    "시험지출처": "source_label",
    "과목명": "course_name_labeled",
    "대분류": "major_category",
    "세부주제": "topic",
    "하위개념": "subtopic",
    "평가항목": "assessment_domain",
    "문항유형": "question_type_labeled",
    "핵심키워드": "concept_tags_raw",
    "교수자확인": "faculty_verified",
    "중복유사문항(2023과만 비교 시) - 문항 유사/주제 유사/중복 없음": "duplicate_similarity_note",
    "애매한점": "ambiguity_note",
    "비고": "ambiguity_note",
}

QUESTION_TYPE_MAP = {
    "단순 개념확인형": "knowledge_recall",
    "임상증례형": "clinical_reasoning",
    "이미지/자료해석형": "image_interpretation",
    "검사해석형": "data_interpretation",
    "치료 선택형": "treatment_selection",
    "치료선택형": "treatment_selection",
    "응급처치형": "emergency_management",
    "가나다라/R형": "complex_multiple_choice",
    "단답형/서술형": "short_answer",
}


@dataclass
class LabelRow:
    row_number: int
    question_number: int
    round_label: str | None
    raw_question_label: str
    labels: dict[str, Any]


def clean_text(value: Any) -> str:
    text = unicodedata.normalize("NFC", str(value or ""))
    text = text.replace("\u00a0", " ")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def normalize_header(value: Any) -> str:
    return clean_text(value)


def normalize_round_label(value: Any) -> str | None:
    text = clean_text(value)
    match = re.search(r"(\d+)\s*차", text)
    if match:
        return f"{match.group(1)}차"
    return None


def normalize_question_number(value: Any) -> int | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        if int(value) == value:
            return int(value)
    match = re.search(r"\d+", clean_text(value))
    return int(match.group(0)) if match else None


def is_subjective_question_label(value: Any) -> bool:
    """Prevent labels like 주1 from being applied to objective question 1."""
    text = clean_text(value)
    return bool(re.match(r"^(주|주관|서술|단답)\s*\d+", text))


def split_concept_tags(value: Any) -> list[str]:
    text = clean_text(value)
    if not text:
        return []
    parts = re.split(r"[,;/·\n]+", text)
    tags: list[str] = []
    seen: set[str] = set()
    for part in parts:
        tag = clean_text(part)
        if not tag or tag in seen:
            continue
        seen.add(tag)
        tags.append(tag)
    return tags


def normalize_question_type(value: Any) -> str | None:
    text = clean_text(value)
    if not text:
        return None
    if "R형" in text or "가나다라" in text:
        if "임상" in text:
            return "complex_clinical_reasoning"
        return "complex_multiple_choice"
    return QUESTION_TYPE_MAP.get(text, re.sub(r"[^0-9A-Za-z가-힣]+", "_", text).strip("_") or None)


def load_label_rows(xlsx_path: Path) -> tuple[list[LabelRow], list[str]]:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to read the labeling workbook.") from exc

    workbook = openpyxl.load_workbook(xlsx_path, data_only=True)
    rows: list[LabelRow] = []
    warnings: list[str] = []

    for worksheet in workbook.worksheets:
        headers = [
            normalize_header(worksheet.cell(1, column).value)
            for column in range(1, worksheet.max_column + 1)
        ]
        header_to_col = {header: idx + 1 for idx, header in enumerate(headers) if header}
        if "문항 번호" not in header_to_col or "시험지출처" not in header_to_col:
            warnings.append(f"{worksheet.title}: required headers not found")
            continue

        for row_idx in range(2, worksheet.max_row + 1):
            raw_question_number = worksheet.cell(row_idx, header_to_col["문항 번호"]).value
            raw_question_label = clean_text(raw_question_number)
            if is_subjective_question_label(raw_question_number):
                continue
            question_number = normalize_question_number(raw_question_number)
            source_value = worksheet.cell(row_idx, header_to_col["시험지출처"]).value
            round_label = normalize_round_label(source_value)
            if question_number is None and not clean_text(source_value):
                continue
            if question_number is None:
                warnings.append(f"{worksheet.title}!row{row_idx}: missing question number")
                continue

            labels: dict[str, Any] = {
                "labeling_status": "labeled",
                "labeling_source": xlsx_path.name,
                "labeling_sheet": worksheet.title,
                "labeling_row": row_idx,
                "source_label": clean_text(source_value),
            }
            for header, key in HEADER_MAP.items():
                col = header_to_col.get(header)
                if not col:
                    continue
                value = clean_text(worksheet.cell(row_idx, col).value)
                if value:
                    labels[key] = value

            raw_type = labels.get("question_type_labeled")
            normalized_type = normalize_question_type(raw_type)
            if normalized_type:
                labels["question_type"] = normalized_type
                labels["question_type_raw"] = raw_type

            tags = split_concept_tags(labels.get("concept_tags_raw"))
            for field in ("major_category", "topic", "subtopic", "assessment_domain"):
                if labels.get(field):
                    tags.append(labels[field])
            deduped_tags: list[str] = []
            seen_tags: set[str] = set()
            for tag in tags:
                if tag and tag not in seen_tags:
                    seen_tags.add(tag)
                    deduped_tags.append(tag)
            labels["concept_tags"] = deduped_tags

            rows.append(
                LabelRow(
                    row_number=row_idx,
                    question_number=question_number,
                    round_label=round_label,
                    raw_question_label=raw_question_label,
                    labels=labels,
                )
            )

    return rows, warnings


def labels_for_record(record: dict[str, Any], label_rows: list[LabelRow]) -> list[LabelRow]:
    round_label = record.get("exam", {}).get("round_label")
    if round_label:
        matching = [row for row in label_rows if row.round_label == round_label]
        if matching:
            return matching
    return label_rows


def apply_labels_to_record(
    record: dict[str, Any],
    label_rows: list[LabelRow],
    *,
    drop_unlabeled: bool,
) -> dict[str, Any]:
    selected_rows = labels_for_record(record, label_rows)
    label_by_question: dict[int, LabelRow] = {}
    duplicate_label_rows: list[dict[str, Any]] = []
    for row in selected_rows:
        if row.question_number in label_by_question:
            duplicate_label_rows.append(
                {
                    "row": row.row_number,
                    "raw_question_label": row.raw_question_label,
                    "source_label": row.labels.get("source_label"),
                    "question_number": row.question_number,
                    "round_label": row.round_label,
                }
            )
            continue
        label_by_question[row.question_number] = row
    matched_questions: list[int] = []
    unmatched_questions: list[int] = []
    questions = []

    for question in record.get("questions", []):
        question_number = normalize_question_number(question.get("question_number"))
        label_row = label_by_question.get(question_number)
        if not label_row:
            unmatched_questions.append(question_number or -1)
            if drop_unlabeled:
                continue
            questions.append(question)
            continue

        merged_labels = {
            **(question.get("labels") if isinstance(question.get("labels"), dict) else {}),
            **label_row.labels,
        }
        question["labels"] = merged_labels
        question["labeling_status"] = "labeled"
        question["review_status"] = question.get("review_status") or "imported"
        question.setdefault("extraction_notes", [])
        if "team_labeling_applied" not in question["extraction_notes"]:
            question["extraction_notes"].append("team_labeling_applied")
        matched_questions.append(question_number or -1)
        questions.append(question)

    matched_set = set(matched_questions)
    unmatched_label_rows = [
        {
            "row": row.row_number,
            "raw_question_label": row.raw_question_label,
            "source_label": row.labels.get("source_label"),
            "question_number": row.question_number,
            "round_label": row.round_label,
        }
        for row in selected_rows
        if row.question_number not in matched_set
    ]

    record["questions"] = questions
    record.setdefault("exam", {})
    record["exam"]["labeling"] = {
        "status": "applied",
        "matched_question_count": len(matched_questions),
        "unmatched_extracted_question_count": len([q for q in unmatched_questions if q != -1]),
        "unmatched_label_row_count": len(unmatched_label_rows),
        "unmatched_label_rows": unmatched_label_rows[:50],
        "duplicate_label_row_count": len(duplicate_label_rows),
        "duplicate_label_rows": duplicate_label_rows[:50],
        "drop_unlabeled": drop_unlabeled,
    }
    return record


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Apply labeling XLSX rows to extracted course exam JSON files.")
    parser.add_argument("json_files", nargs="+", help="Extracted course exam JSON files")
    parser.add_argument("--label-xlsx", required=True, help="Team labeling workbook path")
    parser.add_argument("--output-dir", default="", help="Optional output directory. Defaults to in-place when --in-place is set.")
    parser.add_argument("--in-place", action="store_true", help="Overwrite input JSON files")
    parser.add_argument("--drop-unlabeled", action="store_true", help="Keep only questions that have a matching label row")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    label_rows, warnings = load_label_rows(Path(args.label_xlsx).expanduser())
    if not label_rows:
        print(json.dumps({"error": "no_label_rows", "warnings": warnings}, ensure_ascii=False, indent=2))
        return 1

    output_dir = Path(args.output_dir).expanduser() if args.output_dir else None
    if output_dir:
        output_dir.mkdir(parents=True, exist_ok=True)
    if not args.in_place and output_dir is None:
        print("Use --in-place or --output-dir to choose where enriched JSON should be written.", file=sys.stderr)
        return 2

    summaries = []
    for value in args.json_files:
        path = Path(value).expanduser()
        record = json.loads(path.read_text(encoding="utf-8"))
        enriched = apply_labels_to_record(record, label_rows, drop_unlabeled=args.drop_unlabeled)
        target = path if args.in_place else output_dir / path.name  # type: ignore[operator]
        target.write_text(json.dumps(enriched, ensure_ascii=False, indent=2), encoding="utf-8")
        labeling = enriched.get("exam", {}).get("labeling", {})
        summaries.append(
            {
                "file": target.name,
                "source_exam": enriched.get("exam", {}).get("source_exam"),
                "round_label": enriched.get("exam", {}).get("round_label"),
                "question_count": len(enriched.get("questions", [])),
                "matched_question_count": labeling.get("matched_question_count"),
                "unmatched_label_row_count": labeling.get("unmatched_label_row_count"),
                "duplicate_label_row_count": labeling.get("duplicate_label_row_count"),
                "drop_unlabeled": labeling.get("drop_unlabeled"),
            }
        )

    print(
        json.dumps(
            {
                "label_rows": len(label_rows),
                "warnings": warnings,
                "processed": len(summaries),
                "summaries": summaries,
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
