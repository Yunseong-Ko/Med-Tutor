#!/usr/bin/env python3
"""검토 후 종합 설문(교수님 확정 9문항+자유의견) → 학생별 xlsx 생성 + 배포 폴더 주입.

교수님 검토표와 같은 파일명 규칙(학생_NN_종합설문.xlsx)이라 집계 스크립트가
파일명 번호로 응답자를 식별한다.
"""
import argparse
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path("data_private/professor_items/exports/배포_학생검토")

ITEMS = [
    ("안면타당도", "전체적으로 볼 때, 검토한 문항들은 임상종합평가 문항으로 적절해 보였다."),
    ("내용 포괄성", "검토한 문항들은 의과대학생이 임상종합평가에서 알아야 할 중요한 임상 지식과 내용을 적절히 포함하고 있었다."),
    ("임상적 관련성", "문항에서 제시된 상황과 질문은 실제 임상 상황과 관련성이 높았다."),
    ("임상추론 평가 적절성", "문항들은 단순한 지식 암기뿐 아니라 임상 상황을 해석하고 판단하는 능력을 평가하는 데 적절했다."),
    ("난이도 적절성", "문항의 전반적인 난이도는 의과대학생의 임상종합평가 수준에 적절했다."),
    ("문항 완성도", "문항줄기와 선택지는 전반적으로 명확하고 완성도가 높아 질문의 의도를 이해하기 쉬웠다."),
    ("정답·해설의 교육적 유용성", "제시된 정답과 해설은 정답의 근거를 이해하고 관련 내용을 학습하는 데 도움이 되었다."),
    ("평가도구 활용 가능성", "검토한 문항들은 수정·보완을 거친다면 실제 임상종합평가에 활용할 수 있는 수준이라고 생각한다."),
    ("학습도구 활용 가능성", "이러한 AI 생성 문항은 임상종합평가 준비를 위한 학습 및 자가점검에 도움이 될 것이라고 생각한다."),
]


def make_form(who: str, path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "종합설문"
    ws["A1"] = f"AI 생성 임상종합평가 문항 검토 후 설문 — {who}"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "응답척도: 1=전혀 그렇지 않다 · 2=그렇지 않다 · 3=보통이다 · 4=그렇다 · 5=매우 그렇다"
    ws["A2"].font = Font(size=10, color="666666")
    ws.append([])
    ws.append(["번호", "평가영역", "설문 문항", "응답(1~5)"])
    for c in ws[4]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8F4F4")
    for i, (area, q) in enumerate(ITEMS, 1):
        ws.append([i, area, q, None])
        ws.cell(ws.max_row, 3).alignment = Alignment(wrap_text=True, vertical="top")
    dv = DataValidation(type="list", formula1='"1,2,3,4,5"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"D5:D{4 + len(ITEMS)}")
    r = ws.max_row + 2
    ws.cell(r, 1, "10")
    ws.cell(r, 2, "종합 의견")
    ws.cell(r, 3, "검토한 문항에서 가장 개선이 필요하다고 생각한 점 또는 향후 문항 생성 시 반영하면 좋을 점을 자유롭게 작성해 주십시오.")
    ws.cell(r, 3).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(r + 1, 3, "")   # 자유 기술란
    ws.cell(r + 1, 3).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r + 1].height = 120
    for col, w in (("A", 6), ("B", 22), ("C", 78), ("D", 10)):
        ws.column_dimensions[col].width = w
    wb.save(path)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--students", type=int, default=31)
    args = ap.parse_args()
    made = 0
    for i in range(1, args.students + 1):
        folder = OUT / f"학생_{i:02d}"
        if not folder.exists():
            continue
        make_form(f"학생 {i:02d}", folder / f"학생_{i:02d}_종합설문.xlsx")
        made += 1
    print(f"설문지 {made}개 생성 → 각 학생 폴더")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
