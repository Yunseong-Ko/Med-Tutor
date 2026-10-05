#!/usr/bin/env python3
"""교수 배정표 기준 학생별 검토용 해설집 31세트 생성.

교수님이 배포한 `전체_문항검토_배정표.xlsx`(학생별 배정 시트)의 문항번호를 그대로 따른다.
번호 체계: 1~320 = 1교시 1~80, 2교시 81~160, 3교시 161~240, 4교시 241~320.
학생 검토표(xlsx)의 행 번호와 문서의 문항 번호가 일치해야 하므로 **전역 번호를 표기**한다.

각 문항: 문두 + 검사결과 표 + 이미지 + 선지(정답 표시) + 해설 + 선지별 분석 + 근거
출력: exports/student_review_books/학생_NN_문항해설.docx
"""
import argparse
import json
from pathlib import Path

import openpyxl
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

try:
    from scripts.lab_box_render import add_lab_box
except ImportError:  # 직접 실행(python3 scripts/…) 시
    from lab_box_render import add_lab_box

GEN = Path("data_private/professor_items/generated")
IMG_DIR = Path("data_private/professor_items/images")
# 문서 임베드용 축소본(장변 1600px·JPEG q92 4:4:4). 원본을 그대로 넣으면
# 31개 문서에 이미지가 중복 임베드돼 배포본이 160MB를 넘는다.
DOC_DIR = Path("data_private/course_exams/media/SYNTH_DOC")
PNG_DIR = Path("data_private/course_exams/media/SYNTH_GENERATED")
OUT = Path("data_private/professor_items/exports/student_review_books")


CIRC = {"1": "①", "2": "②", "3": "③", "4": "④", "5": "⑤"}
BOOK_TITLES = {
    "harrison_22e": "Harrison 22e", "sabiston_21e": "Sabiston 21e",
    "nelson_21e": "Nelson 21e", "williams_ob_25e": "Williams Obstetrics 25e",
    "berek_novak_16e": "Berek & Novak 16e", "speroff_9e": "Speroff",
    "kr_psychiatry_3e": "신경정신의학 3판",
}
TEAL = RGBColor(0x0E, 0x7C, 0x7B)
GREY = RGBColor(0x64, 0x74, 0x8B)
RED = RGBColor(0xA1, 0x3F, 0x31)

VERDICT_BADGE = {
    "fully": ("✓ 근거 페이지 검증", TEAL),
    "rerouted": ("◇ 근거 재지정(과별 교과서) — 검증 전", GREY),
    "partially": ("△ 핵심 확인 — 일부 주장은 인용 장 밖(감별질환 등)", RGBColor(0xB4, 0x53, 0x09)),
    "not": ("✗ 인용 장에서 미확인 — 근거 재지정·의학 검수 필요", RED),
}


def resolve_image(name: str):
    if not name:
        return None
    stem = Path(name).stem
    for cand in (DOC_DIR / f"{stem}.jpg", DOC_DIR / f"{stem}.png", PNG_DIR / f"{stem}.png"):
        if cand.exists():
            return cand
    for cand in (PNG_DIR / name, IMG_DIR / name):
        if cand.exists():
            return cand
    return None


def load_items() -> dict:
    """전역 문항번호(1~320) → 문항. 교시 순서대로 이어붙인다."""
    out, n = {}, 0
    for k in range(1, 5):
        for it in json.loads((GEN / f"set_{k}.json").read_text(encoding="utf-8")):
            n += 1
            it = dict(it)
            it["global_no"] = n
            it["period"] = k
            out[n] = it
    return out


def head(doc, text, size=10, color=TEAL):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = color


def body(doc, text, size=9.5, indent=0.4):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(indent)
    p.paragraph_format.space_after = Pt(2)
    p.add_run(str(text)).font.size = Pt(size)


def add_item(doc, it):
    ans = str(it.get("answer", ""))
    p = doc.add_paragraph()
    r = p.add_run(f"문항 {it['global_no']}    정답 {CIRC.get(ans, ans)}")
    r.bold = True
    r.font.size = Pt(12.5)
    meta = " · ".join(x for x in [
        f"{it['period']}교시 {it['no']}번", it.get("subject", ""), it.get("axis", ""),
        it.get("cognitive_level", ""),
        "이미지 문항" if it.get("image") else "",
    ] if x)
    mr = p.add_run(f"    {meta}")
    mr.font.size = Pt(8.5)
    mr.font.color.rgb = GREY

    body(doc, str(it.get("stem", "")), size=9.5, indent=0.2)

    if it.get("lab_box"):
        # 종전엔 검토용 해설집에 검사수치가 빠져 있었다. 설문 T3 반영해 2열 표로
        # 넣고, 파싱 실패 시 원문 그대로 1칸 박스 폴백(유실 없음).
        add_lab_box(doc, it["lab_box"], fallback_size=9.5)

    src = resolve_image(str(it.get("image") or ""))
    if src:
        ip = doc.add_paragraph()
        ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
        ip.paragraph_format.space_after = Pt(2)
        try:
            ip.add_run().add_picture(str(src), width=Cm(9))
        except Exception:
            ip.add_run(f"[이미지: {it.get('image')}]").font.size = Pt(8)

    for k in sorted(it.get("choices", {}), key=lambda x: int(x)):
        pp = doc.add_paragraph()
        pp.paragraph_format.left_indent = Cm(0.6)
        pp.paragraph_format.space_after = Pt(0)
        rr = pp.add_run(f"{CIRC.get(k, k)} {it['choices'][k]}" + (" ◀ 정답" if k == ans else ""))
        rr.font.size = Pt(9.5)
        if k == ans:
            rr.bold = True

    if it.get("explanation"):
        head(doc, "해설")
        body(doc, it["explanation"])

    ce = it.get("choice_explanations") or {}
    if ce:
        head(doc, "선지별 분석")
        for k in sorted(ce, key=lambda x: int(x)):
            row = ce[k] or {}
            txt = str(row.get("why_correct") or row.get("why_attractive") or "").strip()
            if not txt:
                continue
            pp = doc.add_paragraph()
            pp.paragraph_format.left_indent = Cm(0.6)
            pp.paragraph_format.space_after = Pt(1)
            lab = pp.add_run(f"{CIRC.get(k, k)} ")
            lab.bold = True
            lab.font.size = Pt(9)
            if k == ans:
                lab.font.color.rgb = TEAL
            pp.add_run(txt).font.size = Pt(9)

    bits = [f"Harrison 22e Ch.{s.get('chapter')} p.{s.get('printed_page')}"
            for s in (it.get("harrison_sources") or [])]
    bits += [f"{BOOK_TITLES.get(str(t.get('book_id')), t.get('book_id'))} Ch.{t.get('chapter')} (장 포인터)"
             for t in (it.get("textbook_sources") or [])]
    if it.get("disease_concept_id"):
        bits.append(f"온톨로지 개념: {it['disease_concept_id']}")
    if bits:
        head(doc, "근거")
        verdict = str(it.get("entailment_verdict") or "")
        tail_note = "" if verdict else "  (장 포인터 — 문장 수준 검증 전)"
        body(doc, " · ".join(bits) + tail_note)
        badge = VERDICT_BADGE.get(verdict)
        if badge:
            bp = doc.add_paragraph()
            bp.paragraph_format.left_indent = Cm(0.4)
            br = bp.add_run(badge[0])
            br.font.size = Pt(8.5)
            br.bold = True
            br.font.color.rgb = badge[1]
            for pt in (it.get("entailment_unsupported_points") or [])[:3]:
                pp2 = doc.add_paragraph()
                pp2.paragraph_format.left_indent = Cm(0.7)
                rr2 = pp2.add_run(f"· 확인 필요: {pt}")
                rr2.font.size = Pt(8)
                rr2.font.color.rgb = GREY

    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--assign", required=True, help="전체_문항검토_배정표.xlsx")
    args = ap.parse_args()

    items = load_items()
    wb = openpyxl.load_workbook(args.assign)
    ws = wb["학생별 배정"]
    OUT.mkdir(parents=True, exist_ok=True)

    made, missing = 0, set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        student, _, nos = row[0], row[1], row[2]
        if not student or not nos:
            continue
        numbers = [int(x) for x in str(nos).split(",") if str(x).strip()]
        doc = Document()
        for s in doc.sections:
            s.top_margin = s.bottom_margin = Cm(1.8)
            s.left_margin = s.right_margin = Cm(1.8)
        doc.styles["Normal"].font.name = "맑은 고딕"
        doc.styles["Normal"].font.size = Pt(9.5)

        t = doc.add_paragraph()
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        tr = t.add_run(f"AI 생성 필기문항 검토용 해설집 — {student}")
        tr.bold = True
        tr.font.size = Pt(15)
        n = doc.add_paragraph()
        n.alignment = WD_ALIGN_PARAGRAPH.CENTER
        nr = n.add_run(f"배정 {len(numbers)}문항 · 문항 번호는 검토표(엑셀) 행 번호와 같습니다")
        nr.font.size = Pt(9)
        nr.font.color.rgb = GREY
        n2 = doc.add_paragraph()
        n2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        n2r = n2.add_run("의학적 검수 전 AI 생성물 · 외부 배포 금지")
        n2r.font.size = Pt(9)
        n2r.font.color.rgb = RED
        doc.add_paragraph()

        for gid in numbers:
            it = items.get(gid)
            if not it:
                missing.add(gid)
                continue
            add_item(doc, it)

        idx = str(student).replace("학생", "").strip().zfill(2)
        path = OUT / f"학생_{idx}_문항해설.docx"
        doc.save(path)
        made += 1

    print(f"학생별 해설집 {made}개 → {OUT}")
    if missing:
        print(f"  ! 매칭 실패 문항번호 {sorted(missing)[:10]}")
    sizes = sorted(p.stat().st_size for p in OUT.glob("*.docx"))
    if sizes:
        print(f"  파일 크기 {sizes[0]//1024//1024}~{sizes[-1]//1024//1024} MB")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
