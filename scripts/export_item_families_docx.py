#!/usr/bin/env python3
"""문항 패밀리(원본 해설강화 + 유사문항 4) → 검토용 DOCX.

구성: 패밀리별 [원본 문항+강화해설+근거] → [유사문항 1~4 (축회전·CC변형·평행) 각 해설·근거]
이미지 문항은 원본·변형 모두 같은 이미지를 로컬 삽입. 검토용이라 정답·해설 인라인 표시.
입력: generated/mock80_items.json + generated/families_all.json
출력: exports/문항패밀리_검토용.docx
"""
import json
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE = Path("data_private/professor_items")
IMG = BASE / "images"
CIRC = {"1": "①", "2": "②", "3": "③", "4": "④", "5": "⑤"}
TEAL = RGBColor(0x0E, 0x7C, 0x7B)
RED = RGBColor(0xB9, 0x1C, 0x1C)
GRAY = RGBColor(0x64, 0x74, 0x8B)


def add_q(doc, label, stem, choices, answer, explanation, evidence, image=None, kind=""):
    p = doc.add_paragraph()
    r = p.add_run(label + (f" [{kind}]" if kind else ""))
    r.bold = True; r.font.color.rgb = TEAL
    doc.add_paragraph(str(stem or ""))
    if image and (IMG / image).exists():
        ip = doc.add_paragraph(); ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
        try:
            ip.add_run().add_picture(str(IMG / image), width=Cm(8))
        except Exception:
            ip.add_run(f"[이미지: {image}]")
    for k in sorted(choices or {}, key=lambda x: int(x) if str(x).isdigit() else 9):
        if choices[k]:
            doc.add_paragraph(f"{CIRC.get(str(k), k)} {choices[k]}")
    ap = doc.add_paragraph()
    ar = ap.add_run(f"정답 {CIRC.get(str(answer), answer)}")
    ar.bold = True; ar.font.color.rgb = RED
    if explanation:
        ep = doc.add_paragraph(); er = ep.add_run("해설  "); er.bold = True
        ep.add_run(str(explanation))
    if evidence:
        vp = doc.add_paragraph(); vr = vp.add_run(f"근거  {evidence}")
        vr.font.color.rgb = GRAY; vr.font.size = Pt(9)


def main():
    items = {it["no"]: it for it in json.loads((BASE / "generated" / "mock80_items.json").read_text(encoding="utf-8"))}
    fams = json.loads((BASE / "generated" / "families_all.json").read_text(encoding="utf-8"))
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "맑은 고딕"; st.font.size = Pt(10.5)
    h = doc.add_paragraph(); h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hr = h.add_run("문항 패밀리 검토용 — 원본(해설·근거 강화) + 유사문항 4종")
    hr.bold = True; hr.font.size = Pt(15)
    sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("축회전(진단·검사·치료 3축 커버) · C.C변형(정답 변경) · 평행 — 전부 needs_review")
    sr.font.size = Pt(9); sr.font.color.rgb = GRAY
    n_v = 0
    for f in sorted(fams, key=lambda x: x["no"]):
        it = items.get(f["no"])
        if not it:
            continue
        doc.add_paragraph("")
        head = doc.add_paragraph()
        hrr = head.add_run(f"■ 패밀리 {it['no']:02d}  [{it['subject']}]  {it['concept']}")
        hrr.bold = True; hrr.font.size = Pt(12)
        add_q(doc, f"원본 (No {it['no']} · {it['axis']})", it["stem"], it["choices"], it["answer"],
              f.get("enriched_explanation") or it.get("explanation"), f.get("evidence"),
              image=it.get("image"))
        for v in f.get("variants", []):
            n_v += 1
            add_q(doc, f"유사문항 {v['variant']} ({v['axis']})", v["stem"], v["choices"], v["answer"],
                  v.get("explanation"), v.get("evidence"), image=v.get("image"), kind=v.get("kind", ""))
    out = BASE / "exports" / "문항패밀리_검토용.docx"
    doc.save(out)
    print(f"[docx] {out} · 패밀리 {len(fams)} · 유사문항 {n_v}")


if __name__ == "__main__":
    raise SystemExit(main())
