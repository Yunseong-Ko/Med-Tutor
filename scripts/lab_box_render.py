#!/usr/bin/env python3
"""lab_box 문자열 → DOCX 표 렌더 공용 헬퍼 (설문 T3 "검사결과는 표로" 반영).

lab_box는 "검사명 값(참조범위), …" 형태의 한 줄 문자열이다. 참조범위 괄호 안에
쉼표가 있으므로(예: 4,000~10,000) 괄호 깊이 0에서 만나는 구분자만 항목 경계로
인정한다 — '), '로 끊고 ')'를 재부착하는 것과 같은 효과이며, 실데이터에 섞인
'; '·' · '·' / ' 구분과 괄호 없는 항목까지 같은 규칙으로 흡수한다.

정책: 모든 항목이 (검사, 결과, 참조)로 나뉠 때만 표를 만들고, 한 항목이라도
실패하면 표를 포기하고 기존 문장(1칸 박스) 렌더로 폴백한다 — 내용 유실 금지.
lab_box가 없거나 빈 기존 문항은 아무 영향도 받지 않는다(하위호환).
"""
import re

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

# 수치가 없는 정성 결과로 인정하는 꼬리 어절(startswith 매칭)
_QUAL_TAILS = (
    "음성", "양성", "강양성", "약양성", "없음", "다수", "정상",
    "증가", "감소", "상승", "저하", "미검출", "무반응", "폐쇄",
)
# 값 어절의 숫자 머리(부등호·부호 허용). 숫자 뒤 한글은 단위 글자만 허용해
# '25-수산화비타민D'(검사명) 같은 어절을 값으로 오인하지 않는다.
_NUM_HEAD = re.compile(r"^[<>≥≤≈±+]?\d[\d,.]*")
_UNIT_HANGUL = set("만천억초분회일")
_HANGUL = re.compile(r"[가-힣]")


def _depth0_split(text, delims):
    """괄호 밖(depth 0)에서 만나는 delims로만 분리 — 참조범위 안의 쉼표 보호."""
    out, buf, depth, i = [], [], 0, 0
    while i < len(text):
        ch = text[i]
        if ch == "(":
            depth += 1
        elif ch == ")":
            depth = max(0, depth - 1)
        if depth == 0:
            hit = next((d for d in delims if text.startswith(d, i)), None)
            if hit:
                out.append("".join(buf).strip())
                buf = []
                i += len(hit)
                continue
        buf.append(ch)
        i += 1
    out.append("".join(buf).strip())
    return [e for e in out if e]


def _is_value_token(tok):
    m = _NUM_HEAD.match(tok)
    if not m:
        return False
    return all(c in _UNIT_HANGUL for c in _HANGUL.findall(tok[m.end():]))


def _split_entry(entry):
    """'검사명 값(참조범위)' 한 항목 → (검사, 결과, 참조). 실패 시 None."""
    entry = entry.strip()
    ref = ""
    if entry.endswith(")"):
        idx = entry.rfind("(")
        if idx > 0:
            ref = entry[idx + 1:-1].strip()
            entry = entry[:idx].strip()
    if ": " in entry:  # '자궁경부 세포검사: 고등급 편평상피내병변' 서술형
        name, val = (s.strip() for s in entry.split(": ", 1))
        return (name, val, ref) if name and val else None
    toks = entry.split()
    for i in range(1, len(toks)):  # i>0: 검사명이 비지 않게
        if _is_value_token(toks[i]):
            return (" ".join(toks[:i]), " ".join(toks[i:]), ref)
    if len(toks) >= 2 and toks[-1].startswith(_QUAL_TAILS):
        return (" ".join(toks[:-1]), toks[-1], ref)
    return None


def parse_lab_box(text):
    """lab_box 전체 → [(검사, 결과, 참조범위)]. 표로 만들 수 없으면 빈 리스트."""
    s = str(text or "").strip()
    if not s:
        return []
    entries = _depth0_split(s, (", ", "; ", " · "))
    if len(entries) == 1:  # 'A 1(…) / B 2(…)'처럼 슬래시로만 나열한 lab_box
        entries = _depth0_split(s, (", ", "; ", " · ", " / "))
    if len(entries) < 2:  # 항목 1개는 표의 이득이 없다 — 문장 유지
        return []
    rows = []
    for e in entries:
        parsed = _split_entry(e)
        if parsed is None:
            return []
        rows.append(parsed)
    return rows


def _fill(cell, text, size, bold=False):
    """셀 텍스트 채우기 + 좁은 셀 여백(상하 40·좌우 80 dxa)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, w in (("top", "40"), ("left", "80"), ("bottom", "40"), ("right", "80")):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), w)
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tc_pr.append(mar)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold


def add_lab_table(doc, lab_box, font_size=8.5):
    """파싱되면 2열 표(검사 | 결과(참조치))를 추가하고 True, 아니면 False."""
    rows = parse_lab_box(lab_box)
    if not rows:
        return False
    has_ref = any(ref for _, _, ref in rows)
    widths = (Cm(4.6), Cm(9.4))
    t = doc.add_table(rows=len(rows) + 1, cols=2)
    t.style = "Table Grid"
    t.autofit = False
    for col, w in zip(t.columns, widths):
        col.width = w
    header = ("검사", "결과 (참조치)" if has_ref else "결과")
    for r, (name, result) in enumerate([header] + [
        (name, f"{val} ({ref})" if ref else val) for name, val, ref in rows
    ]):
        cells = t.rows[r].cells
        cells[0].width, cells[1].width = widths
        _fill(cells[0], name, font_size, bold=(r == 0))
        _fill(cells[1], result, font_size, bold=(r == 0))
    return True


def add_lab_box(doc, lab_box, font_size=8.5, fallback_size=10):
    """표 파싱 성공 시 2열 표, 실패 시 기존 1칸 박스(문장 그대로) — 유실 없음.

    반환값 'table'|'box'로 호출부·검증 스크립트가 파싱 성공률을 집계할 수 있다.
    """
    if add_lab_table(doc, lab_box, font_size=font_size):
        return "table"
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.rows[0].cells[0].paragraphs[0].add_run(str(lab_box)).font.size = Pt(fallback_size)
    return "box"
