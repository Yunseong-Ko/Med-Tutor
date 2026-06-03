#!/usr/bin/env python3
"""
Parse KAMC-style answer key XLSX files into period/question answer maps.
"""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

import openpyxl


CIRCLE_TO_STR = {
    "①": "1",
    "②": "2",
    "③": "3",
    "④": "4",
    "⑤": "5",
}


def normalize_answer(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    if text in CIRCLE_TO_STR:
        return CIRCLE_TO_STR[text]
    match = re.search(r"[①②③④⑤1-5]", text)
    if not match:
        return None
    return CIRCLE_TO_STR.get(match.group(0), match.group(0))


def normalize_period(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    match = re.search(r"(\d+)\s*교시", text)
    if not match:
        return None
    return f"{match.group(1)}교시"


def parse_answer_key_xlsx(path: Path) -> dict:
    workbook = openpyxl.load_workbook(path, data_only=True)
    rows: list[dict] = []
    by_period: dict[str, dict[str, str]] = {}

    for sheet in workbook.worksheets:
        for row_index in range(1, sheet.max_row + 1):
            first = sheet.cell(row_index, 1).value
            if str(first).strip() != "문항번호":
                continue

            question_numbers: dict[int, int] = {}
            for col_index in range(2, sheet.max_column + 1):
                raw_number = sheet.cell(row_index, col_index).value
                if raw_number is None:
                    continue
                try:
                    question_numbers[col_index] = int(raw_number)
                except (TypeError, ValueError):
                    continue

            for answer_row in range(row_index + 1, min(row_index + 5, sheet.max_row) + 1):
                period = normalize_period(sheet.cell(answer_row, 1).value)
                if not period:
                    continue
                for col_index, question_number in question_numbers.items():
                    answer = normalize_answer(sheet.cell(answer_row, col_index).value)
                    if not answer:
                        continue
                    by_period.setdefault(period, {})[str(question_number)] = answer
                    rows.append(
                        {
                            "sheet": sheet.title,
                            "period_label": period,
                            "question_number": question_number,
                            "answer": answer,
                        }
                    )

    return {
        "source_file": path.name,
        "periods": sorted(by_period),
        "question_count_by_period": {
            period: len(answers)
            for period, answers in sorted(by_period.items())
        },
        "answers_by_period": by_period,
        "rows": rows,
    }


def apply_answer_key_to_record(record: dict, answer_key: dict, *, period_label: str | None = None) -> dict:
    exam = record.get("exam", {})
    selected_period = period_label or exam.get("period_label")
    answers_by_period = answer_key.get("answers_by_period", {})

    if selected_period and selected_period in answers_by_period:
        answer_map = answers_by_period[selected_period]
    elif len(answers_by_period) == 1:
        selected_period, answer_map = next(iter(answers_by_period.items()))
    else:
        answer_map = {}

    matched = 0
    mismatched: list[int] = []
    missing: list[int] = []
    for question in record.get("questions", []):
        key = str(question.get("question_number"))
        answer = answer_map.get(key)
        if not answer:
            missing.append(question.get("question_number"))
            continue
        existing = str(question.get("answer") or "").strip()
        if existing and existing != answer:
            mismatched.append(question.get("question_number"))
            question.setdefault("extraction_notes", []).append(
                f"answer_overridden_by_xlsx:{existing}->{answer}"
            )
        question["answer"] = answer
        question["answer_source"] = "answer_key_xlsx"
        question["needs_review"] = bool(question.get("review_reasons"))
        if "answer_not_extracted" in question.get("review_reasons", []):
            question["review_reasons"] = [
                reason
                for reason in question["review_reasons"]
                if reason != "answer_not_extracted"
            ]
            question["needs_review"] = bool(question["review_reasons"])
        matched += 1

    record["answer_key"] = {
        "source_file": answer_key.get("source_file"),
        "period_label": selected_period,
        "matched_count": matched,
        "missing_question_numbers": missing,
        "mismatched_question_numbers": mismatched,
        "available_periods": answer_key.get("periods", []),
    }
    return record


def safe_summary(answer_key: dict) -> dict:
    return {
        "source_file": answer_key.get("source_file"),
        "periods": answer_key.get("periods", []),
        "question_count_by_period": answer_key.get("question_count_by_period", {}),
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Parse answer key XLSX into a safe JSON summary.")
    parser.add_argument("input", help="Answer key XLSX path")
    parser.add_argument("--output", help="Optional JSON output path")
    parser.add_argument("--full", action="store_true", help="Include full answer map in stdout/output")
    args = parser.parse_args()

    answer_key = parse_answer_key_xlsx(Path(args.input).expanduser())
    payload = answer_key if args.full else safe_summary(answer_key)
    if args.output:
        Path(args.output).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
