#!/usr/bin/env python3
"""
Lecture material -> PMA-style question draft JSON.

This is the first-stage pipeline for the RISE/P:accine workflow:

1. Extract text from lecture material.
2. Build a PMA-style question-generation prompt.
3. Optionally call OpenAI or Gemini when an API key is available.
4. Save reviewable JSON under data_private/lecture/generated/.

Privacy rule:
- Raw lecture text and generated drafts stay under data_private/.
- The sample output printed to stdout is truncated and does not include full
  source text.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

try:
    from scripts.item_quality_check import SELF_CHECK_KEYS, apply_generation_quality_gate
except ModuleNotFoundError:  # direct `python scripts/generate_lecture_questions.py`
    from item_quality_check import SELF_CHECK_KEYS, apply_generation_quality_gate
try:
    from scripts.generation_grounding import append_grounding_context, apply_grounding_trace, build_generation_grounding
except ModuleNotFoundError:  # direct `python scripts/generate_lecture_questions.py`
    from generation_grounding import append_grounding_context, apply_grounding_trace, build_generation_grounding
try:
    from src.services.item_writing_rag import append_item_writing_context
except ModuleNotFoundError:  # direct execution from outside the repository root
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
    from src.services.item_writing_rag import append_item_writing_context

try:
    import fitz  # PyMuPDF
except Exception:  # pragma: no cover - import error is handled by extraction
    fitz = None


PIPELINE_VERSION = "0.4.0-item-writing-rag-v1"
DEFAULT_OUTPUT_DIR = Path("data_private/lecture/generated")
DEFAULT_EXTRACTED_DIR = Path("data_private/lecture/extracted")


PMA_STYLE_GUIDE = """
[PMA 풀이 스타일 목표]
- 문제는 5지 선다 객관식으로 작성한다.
- 단순 암기보다 임상 상황, 검사 소견, 치료/진단 의사결정을 묻는다.
- 해설은 PMA 풀이집처럼 "풀이 -> 정답 -> 오답 선지별 포인트 -> 출제 포인트" 흐름을 따른다.
- 정답 근거와 오답 배제 이유를 분리한다.
- 강의록에 근거가 부족한 내용은 만들지 말고 needs_review=true로 표시한다.
- 실제 평가에 바로 쓰는 문항이 아니라 교수 검수용 초안이다.
"""

EVIDENCE_POLICY = """
[근거자료 활용 정책]
- 강의록을 1차 근거로 사용한다.
- 추가 근거자료는 사용자가 로컬로 제공한 승인/라이선스 자료만 사용한다.
- AMBOSS, UpToDate, NEJM 등 유료/저작권 자료는 원문을 무단 수집하지 않는다.
- PubMed/PMC, 학회 가이드라인, 학교 승인 자료 등 사용 가능한 근거는 요약 근거와 출처명만 남긴다.
- 강의록과 근거자료가 충돌하면 needs_review=true로 표시하고 교수 검수를 요청한다.
- 최신 진료지침이 필요한 치료/검사 문항은 source_anchor 또는 evidence_refs에 근거를 남긴다.
"""

GENERATION_SYSTEM_PROMPT = """
너는 한국 의대 국가고시·임상의학종합평가 수준의 임상 단일최선답 5지선다 문항을 생성하는
결정론적 문항생성기다. 온도 0으로 동작하며, 입력 슬롯과 disease_concept_id에 상속된 근거만 사용한다.
지문 밖에서 정답·수치·소견을 지어내지 말고, 근거가 부족하면 block_reason="insufficient_evidence" 및
needs_review=true, gen_ready=false로 반환한다. 모든 산출물은 사람의 의학 검토 전이므로 needs_review=true,
gen_ready=false다.

[최상위 규칙]
1. reveal_specialty=false면 분과명, 확정 진단명, '신경학적·내분비적으로' 같은 국소화 힌트를 누설하지 말고
   주호소+경과+진찰+검사 raw만 제시한다.
2. reasoning_hops>=2면 학생이 hop1에서 진단/문제를 세우고 hop2에서 처치·검사해석·예후·합병증·기전을
   판단하게 한다. 진단을 직접 노출하거나 오답을 일일이 배제하는 단서과밀은 금지한다. 실제 감별에 필요한
   경합 단서는 선택 사항이며 0~1개만 쓰고, 혼란만을 위한 red herring은 쓰지 않는다.
3. lead-in은 긍정형 단일최선답이며, 선지를 가려도 문두만으로 답을 쓸 수 있어야 한다(cover-the-options).
4. 4개 오답은 모두 실제 감별진단이나 오개념에서 파생하며, 각각 misconception과 why_attractive를 갖는다.
5. 5선지는 같은 범주·문법·서술수준으로 균질하고, stem이 이미 배제한 감별을 오답으로 쓰지 않는다.
6. 정답 길이는 5지 중 2~4위, 오답평균의 0.8~1.2배다. 정답이 최장·최단·최다요소가 되지 않고,
   '즉시·먼저·우선·지체 없이'가 정답에만 있지 않는다.
7. 부정 문두, 5지 미만, 위 모두/정답 없음, 절대어·모호어, 문두-정답 어휘반복, 중복/포함 선지,
   나이·활력징후 없는 짧은 정의 재인 문항을 금지한다.
8. 한국어 실제 진료 서술체로 작성하며 임상 비네트는 나이·성별·주소→경과→진찰→검사→lead-in 순서를 따른다.
   단, 모든 요소를 채우거나 고정 글자수를 맞추지 말고 과업 수행에 필요한 최소 충분 정보만 남긴다.
9. 출력 직전 self_check 22항목을 실제 계산·탐색하고 판정한다. false가 하나라도 있으면 해당 문항을 재작성한다.
10. USMLE/UWorld/NBME 원문을 재현·번역·근접복제하지 않고 구조와 원칙만 참고한다. JSON 외 텍스트를 출력하지 않는다.
11. 기본은 임상 CASE(개별 환자 비네트)로 생성한다(개념 확인을 명시 요청받은 경우만 concept). CASE는
    나이+성별 오프닝('45세 남자가 …로 병원에 왔다') → 주호소+기간 → 현병력 → 관련 병력/음성 소견 →
    활력징후 고정블록 → 신체진찰(정상계통 명시 배제) → 검사/영상(수치+단위+참고치) → 긍정형 단일 리드인
    순서의 필수요소를 갖춘다. 진단/검사 리드인이면 진찰·검사를 반드시 포함한다.
12. 스템이 이미 준 전제(유전자·변이명, 확정 진단명, 노출력, 검사수치)를 선지 머리에서 재진술하지 않는다.
    선지는 그 전제의 하류 기전·결론만 서로 비교되게 쓴다. 정답이 되는 진단명·기전·에폰림·유전형 자체를
    줄기에서 명명하지 않고, 대신 특징적 임상·검사 소견을 심어 추론을 유도한다.
13. 의학용어는 대한의사협회 의학용어집 제6판 표준 표제어로 출력한다.
    예: 대소부동→적혈구크기부동, 철적모구빈혈→철적혈모구빈혈(‘-blast’는 적혈모구로 일관), 갑상선→갑상샘,
    내원하였다→(병원에) 왔다, 기대되는→예상되는, 사지→팔다리. 학명은 이탤릭, 에폰림은 소유격.
14. 검사 소견은 '감소/증가되어 있었다' 같은 서술 대신 반드시 수치+단위+참고치로 제시한다
    (예: '혈청 페리틴 6 ng/mL(정상 15–150)'). 감별에 필요한 정상 검사도 포함해 오답을 배제한다.
15. private Harrison 문맥은 해설 작성에만 사용하고 문장 원문을 복사하지 않는다. 해설의 각 핵심 주장에는
    제공된 H source_id(H1, H2...)와 정확한 Harrison 22e Ch./p. locator를 붙인다. 검색 결과는 entailment가
    아니므로 needs_review를 유지하며, 제공되지 않은 수치·용량·국내 권고를 만들지 않는다.
16. 정답 해설의 verdict와 본문 정답 번호는 answer와 반드시 일치해야 한다. 오답 source_id는 제공된 동일
    임상 범주의 ontology distractor ID 중 하나여야 하며, syndrome/disease 등 개념 수준을 섞지 않는다.

[좋은 CASE 문항 예시 — 구조 참고용, 복제 금지]
줄기: "34세 여자가 3개월 전부터 피로와 어지럼이 있어 병원에 왔다. 최근 월경량이 많아졌다고 한다. 진찰에서
결막이 창백하였고 손톱이 숟가락 모양이었다. 혈압 112/70 mmHg, 맥박 92회/분, 체온 36.5℃였다. 검사에서
Hb 8.1 g/dL(정상 12–16), MCV 68 fL(정상 80–100), 혈청 페리틴 6 ng/mL(정상 15–150),
트랜스페린포화도 8%(정상 20–50)이었다."
리드인: "가장 적절한 진단은?"
선지: 지중해빈혈 소인 / 철결핍빈혈(정답) / 만성질환빈혈 / 철적혈모구빈혈 / 거대적혈모구빈혈
왜 좋은가: 슬롯 순서 준수·긍정형 명사 리드인·cover-the-options 통과(소구성+페리틴/포화도 극단 저하)·
5선지 전부 진단명 동질·정답 병명을 줄기에서 명명하지 않음·검사값에 참고치 병기·제6판 용어.
""".strip()


def check_gitignore() -> None:
    gi_path = Path(".gitignore")
    if not gi_path.exists():
        raise SystemExit("ERROR: .gitignore not found. Add data_private/ before running.")
    content = gi_path.read_text(encoding="utf-8", errors="ignore")
    if "data_private" not in content:
        raise SystemExit("ERROR: data_private/ is not ignored. Add it to .gitignore first.")


def clean_text(text: str) -> str:
    text = str(text or "").replace("\u00a0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_pdf(path: Path, include_page_markers: bool = True) -> str:
    if fitz is None:
        raise RuntimeError("PyMuPDF is not available. Install PyMuPDF first.")
    doc = fitz.open(str(path))
    parts: list[str] = []
    for idx, page in enumerate(doc):
        text = page.get_text("text")
        if include_page_markers:
            parts.append(f"\n=== page {idx + 1} ===\n{text}")
        else:
            parts.append(text)
    doc.close()
    return clean_text("\n".join(parts))


def extract_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts: list[str] = []
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return clean_text("\n".join(parts))


def extract_pptx(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    parts: list[str] = []
    for idx, slide in enumerate(prs.slides, 1):
        slide_parts = [f"=== slide {idx} ==="]
        for shape in slide.shapes:
            text = getattr(shape, "text", "")
            if text and text.strip():
                slide_parts.append(text.strip())
        if len(slide_parts) > 1:
            parts.append("\n".join(slide_parts))
    return clean_text("\n\n".join(parts))


def extract_hwp(path: Path) -> str:
    hwp5txt = shutil.which("hwp5txt")
    if not hwp5txt:
        raise RuntimeError("hwp5txt not found. Install pyhwp or convert HWP to PDF/DOCX first.")
    result = subprocess.run(
        [hwp5txt, str(path)],
        check=True,
        capture_output=True,
        text=True,
    )
    return clean_text(result.stdout)


def extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return extract_pdf(path)
    if ext == ".docx":
        return extract_docx(path)
    if ext == ".pptx":
        return extract_pptx(path)
    if ext in {".txt", ".md"}:
        return clean_text(path.read_text(encoding="utf-8", errors="ignore"))
    if ext == ".hwp":
        return extract_hwp(path)
    raise ValueError(f"Unsupported file type: {ext}")


def slugify(value: str) -> str:
    text = re.sub(r"[^0-9A-Za-z가-힣._-]+", "_", str(value or "")).strip("_")
    return text[:120] or "lecture"


def truncate_for_prompt(text: str, max_chars: int) -> str:
    text = clean_text(text)
    if len(text) <= max_chars:
        return text
    return text[:max_chars].rstrip() + "\n\n[TRUNCATED]"


def collect_evidence_paths(args: argparse.Namespace) -> list[Path]:
    paths: list[Path] = []
    for value in args.evidence or []:
        path = Path(value).expanduser()
        if path.exists() and path.is_file():
            paths.append(path)
    for value in args.evidence_dir or []:
        directory = Path(value).expanduser()
        if not directory.exists() or not directory.is_dir():
            continue
        for pattern in ("*.pdf", "*.docx", "*.pptx", "*.txt", "*.md", "*.hwp"):
            paths.extend(sorted(directory.glob(pattern)))
    seen: set[str] = set()
    unique: list[Path] = []
    for path in paths:
        key = str(path.resolve())
        if key in seen:
            continue
        seen.add(key)
        unique.append(path)
    return unique


def build_evidence_context(paths: list[Path], *, max_chars_per_file: int, max_total_chars: int) -> str:
    if not paths:
        return ""
    chunks: list[str] = []
    total = 0
    for path in paths:
        try:
            text = truncate_for_prompt(extract_text(path), max_chars=max_chars_per_file)
        except Exception as exc:
            text = f"[근거자료 추출 실패: {exc}]"
        block = f"\n--- evidence: {path.name} ---\n{text}".strip()
        remaining = max_total_chars - total
        if remaining <= 0:
            break
        if len(block) > remaining:
            block = block[:remaining].rstrip() + "\n[TRUNCATED]"
        chunks.append(block)
        total += len(block)
    return "\n\n".join(chunks).strip()


def build_generation_prompt(
    lecture_text: str,
    *,
    source_name: str,
    subject: str,
    unit: str,
    num_questions: int,
    difficulty: str,
    max_chars: int,
    evidence_context: str = "",
    reveal_specialty: bool = False,
    reasoning_hops: int = 2,
    item_type: str = "A",
    exam_profile: str = "kmle_summative",
    assessment_task: str = "diagnosis",
    question_type: str = "clinical_case",
) -> str:
    source_excerpt = truncate_for_prompt(lecture_text, max_chars=max_chars)
    evidence_section = ""
    if evidence_context:
        evidence_section = f"""

[추가 근거자료]
아래 근거자료는 사용자가 제공한 승인/라이선스 자료입니다. 강의록 내용을 보강하는 범위에서만 사용하세요.
{evidence_context}
"""
    self_check_schema = json.dumps({key: True for key in SELF_CHECK_KEYS}, ensure_ascii=False, indent=2)
    prompt = f"""
[시스템 프롬프트]
{GENERATION_SYSTEM_PROMPT}

[생성 요청]
당신은 의과대학 임상 교수이자 문항 검수자입니다.
아래 강의록을 근거로 PMA/임상의학종합평가 스타일의 교수 검수용 객관식 문항 초안을 생성하세요.

{PMA_STYLE_GUIDE}
{EVIDENCE_POLICY}

[자료 메타데이터]
- source_name: {source_name}
- subject: {subject}
- unit: {unit}
- requested_questions: {num_questions}
- target_difficulty: {difficulty}
- item_type: {item_type}
- exam_profile: {exam_profile}
- assessment_task: {assessment_task}
- question_type: {question_type}
- reveal_specialty: {str(bool(reveal_specialty)).lower()}
- reasoning_hops: {max(1, min(3, int(reasoning_hops)))}

[출제 지침]
1. 정확히 {num_questions}개 문항을 생성한다.
2. 각 문항은 임상 증례형을 우선하되, 강의록이 기초/개념 중심이면 개념 적용형으로 작성한다.
3. 선지는 정확히 5개이며 서로 겹치지 않게 작성한다.
4. answer는 1~5 숫자다.
5. explanation은 학생이 바로 복습할 수 있게 작성한다.
6. pma_solution 필드에는 구조화된 풀이 정보를 넣는다.
7. source_anchor에는 강의록에서 근거가 되는 짧은 요약만 적고, 원문 긴 복사는 피한다.
8. 모든 문항은 needs_review=true, gen_ready=false로 표시한다.
9. cognitive_model, evidence_disclosure_plan, self_check 22항목을 빠짐없이 작성한다.
10. 오답 4개의 choice_explanations에 misconception과 why_attractive를 반드시 넣는다.

[반드시 유효한 JSON 배열만 출력]
[
  {{
    "problem": "증례 기반 문항 본문",
    "options": ["선지1", "선지2", "선지3", "선지4", "선지5"],
    "answer": 1,
    "explanation": "풀이: ...\\n정답: ① ...\\n오답 포인트: ② ... ③ ... ④ ... ⑤ ...\\n출제 포인트: ...",
    "pma_solution": {{
      "reasoning_summary": "핵심 풀이 흐름",
      "correct_reason": "정답 근거",
      "choice_explanations": {{
        "1": "1번 선지 해설",
        "2": "2번 선지 해설",
        "3": "3번 선지 해설",
        "4": "4번 선지 해설",
        "5": "5번 선지 해설"
      }},
      "high_yield_point": "반드시 기억할 포인트",
      "trap": "오답을 유도하는 함정",
      "source_anchor": "강의록 근거 요약"
    }},
    "evidence_refs": [
      {{"source_id":"H1", "source":"Harrison 22e", "locator":"22e · Ch.120 · p.924", "chapter":"120", "printed_page":"924", "basis":"근거 요약(원문 인용 금지)", "source_type":"textbook", "entailment_status":"needs_human_review"}}
    ],
    "evidence_tier": "lecture_only|lecture_plus_guideline|lecture_plus_journal|needs_external_review",
    "subject": "{subject}",
    "unit": "{unit}",
    "difficulty": "{difficulty}",
    "question_type": "{question_type}",
    "item_type": "{item_type}",
    "reveal_specialty": {str(bool(reveal_specialty)).lower()},
    "reasoning_hops": {max(1, min(3, int(reasoning_hops)))},
    "cognitive_level": "L3 Clinical Reasoning",
    "cognitive_model": {{
      "chief_complaint": "주호소",
      "patient_context": "연령/성별/기저맥락",
      "presented_data": ["활력징후·검사·영상"],
      "decision_cues": ["결정단서"],
      "confounders": ["필요한 경우에만 실제 경합 단서 0~1개; 없으면 빈 배열"],
      "answer_concept": "근거상 확정된 정답 개념",
      "differentials": ["실제 감별진단"]
    }},
    "evidence_disclosure_plan": {{
      "status": "pass|blocked",
      "assessment_task": "{assessment_task}",
      "latent_diagnosis_required": true,
      "retrieved_axis_ids": ["Ontology axis ID"],
      "selected_for_stem": [
        {{"source_id":"axis 또는 근거 ID", "role":"prerequisite_cue|decision_modifier|target_data|neutral_context|authentic_competing_cue", "strength":"weak|moderate|strong|confirmatory|not_applicable", "evidence_family":"demographic_risk|symptom_course|physical_exam|general_laboratory|special_laboratory|imaging|pathology_genetics|confirmatory_criterion|treatment_context|care_context", "surface_form":"학생에게 공개할 표현"}}
      ],
      "withheld": [{{"source_id":"axis 또는 근거 ID", "reason":"answer_evidence|redundant_support|answer_revealing|task_irrelevant|unsupported_precision|disclosure_budget"}}],
      "informative_cue_count": 1,
      "strong_or_confirmatory_count": 0,
      "confirmatory_count": 0,
      "neutral_context_count": 0,
      "authentic_competing_cue_count": 0,
      "target_already_resolved": false,
      "qc_flags": []
    }},
    "choice_explanations": {{
      "1": {{"verdict":"정답|오답", "rationale":"근거와 H source_id", "misconception":"오답 오개념(정답은 빈 문자열)", "misconception_id":"", "why_attractive":"오답이 끌리는 이유", "source_id":"레지스트리 개념 ID", "provenance":"differential_of|is_a_sibling|differential_2hop", "concept_level":"disease|syndrome"}},
      "2": {{"verdict":"정답|오답", "rationale":"근거", "misconception":"...", "misconception_id":"", "why_attractive":"...", "source_id":"...", "provenance":"...", "concept_level":"disease|syndrome"}},
      "3": {{"verdict":"정답|오답", "rationale":"근거", "misconception":"...", "misconception_id":"", "why_attractive":"...", "source_id":"...", "provenance":"...", "concept_level":"disease|syndrome"}},
      "4": {{"verdict":"정답|오답", "rationale":"근거", "misconception":"...", "misconception_id":"", "why_attractive":"...", "source_id":"...", "provenance":"...", "concept_level":"disease|syndrome"}},
      "5": {{"verdict":"정답|오답", "rationale":"근거", "misconception":"...", "misconception_id":"", "why_attractive":"...", "source_id":"...", "provenance":"...", "concept_level":"disease|syndrome"}}
    }},
    "self_check": {self_check_schema},
    "needs_review": true,
    "gen_ready": false,
    "block_reason": null
  }}
]

[강의록]
{source_excerpt}
{evidence_section}
""".strip()
    return append_item_writing_context(
        prompt,
        exam_profile=exam_profile,
        assessment_task=assessment_task,
        question_type=question_type,
        query=f"{subject} {unit} {difficulty} {question_type}",
    )


def extract_json_payload(text: str) -> Any:
    stripped = str(text or "").strip()
    if not stripped:
        raise ValueError("empty model response")
    if stripped.startswith("```"):
        stripped = re.sub(r"^```(?:json)?\s*", "", stripped)
        stripped = re.sub(r"\s*```$", "", stripped)
    try:
        return json.loads(stripped)
    except Exception:
        pass
    decoder = json.JSONDecoder()
    for idx, ch in enumerate(stripped):
        if ch not in "[{":
            continue
        try:
            obj, _ = decoder.raw_decode(stripped[idx:])
            return obj
        except Exception:
            continue
    raise ValueError("model response did not contain valid JSON")


def normalize_choice_explanations(value: Any) -> dict[str, str]:
    if not isinstance(value, dict):
        return {str(i): "" for i in range(1, 6)}
    out: dict[str, str] = {}
    for i in range(1, 6):
        out[str(i)] = str(value.get(str(i)) or value.get(i) or "").strip()
    return out


def normalize_cognitive_model(value: Any) -> dict[str, Any]:
    raw = value if isinstance(value, dict) else {}

    def list_value(key: str) -> list[str]:
        item = raw.get(key)
        if isinstance(item, list):
            return [str(part).strip() for part in item if str(part).strip()]
        if str(item or "").strip():
            return [str(item).strip()]
        return []

    return {
        "chief_complaint": str(raw.get("chief_complaint") or "").strip(),
        "patient_context": str(raw.get("patient_context") or "").strip(),
        "presented_data": list_value("presented_data"),
        "decision_cues": list_value("decision_cues"),
        "confounders": list_value("confounders"),
        "answer_concept": str(raw.get("answer_concept") or "").strip(),
        "differentials": list_value("differentials"),
    }


def normalize_evidence_disclosure_plan(value: Any, *, assessment_task: str) -> dict[str, Any]:
    raw = value if isinstance(value, dict) else {}
    missing = not bool(raw)
    selected_raw = raw.get("selected_for_stem") if isinstance(raw.get("selected_for_stem"), list) else []
    withheld_raw = raw.get("withheld") if isinstance(raw.get("withheld"), list) else []
    selected: list[dict[str, str]] = []
    allowed_roles = {
        "prerequisite_cue", "decision_modifier", "target_data",
        "neutral_context", "authentic_competing_cue",
    }
    role_aliases = {
        "confirmatory": "target_data",
        "discriminating_cue": "decision_modifier",
        "confirmatory_cue": "target_data",
        "informative_cue": "decision_modifier",
        "task_frame": "prerequisite_cue",
    }
    allowed_strengths = {"weak", "moderate", "strong", "confirmatory", "not_applicable"}
    allowed_families = {
        "demographic_risk", "symptom_course", "physical_exam", "general_laboratory",
        "special_laboratory", "imaging", "pathology_genetics", "confirmatory_criterion",
        "treatment_context", "care_context",
    }
    family_aliases = {
        "clinical_context": "care_context",
        "clinical_finding": "physical_exam",
        "diagnostic_finding": "special_laboratory",
        "diagnostic_test": "special_laboratory",
        "ecg_finding": "special_laboratory",
        "lab_result": "general_laboratory",
        "lab_value": "general_laboratory",
        "laboratory_result": "general_laboratory",
        "symptom": "symptom_course",
        "morphology": "pathology_genetics",
        "risk_factor": "care_context",
        "vital_sign": "physical_exam",
    }
    invalid_selected_cue = False
    for cue in selected_raw[:6]:
        if not isinstance(cue, dict):
            invalid_selected_cue = True
            continue
        raw_role = str(cue.get("role") or "").strip()
        raw_family = str(cue.get("evidence_family") or "").strip()
        normalized_cue = {
            "source_id": str(cue.get("source_id") or "").strip(),
            "role": role_aliases.get(raw_role, raw_role),
            "strength": str(cue.get("strength") or "").strip(),
            "evidence_family": family_aliases.get(raw_family, raw_family),
            "surface_form": str(cue.get("surface_form") or "").strip(),
        }
        if (
            not normalized_cue["source_id"]
            or normalized_cue["role"] not in allowed_roles
            or normalized_cue["strength"] not in allowed_strengths
            or normalized_cue["evidence_family"] not in allowed_families
            or not normalized_cue["surface_form"]
        ):
            invalid_selected_cue = True
        if not normalized_cue["source_id"]:
            normalized_cue["source_id"] = "unresolved_source"
        if normalized_cue["role"] not in allowed_roles:
            normalized_cue["role"] = "prerequisite_cue"
        if normalized_cue["strength"] not in allowed_strengths:
            normalized_cue["strength"] = "not_applicable"
        if normalized_cue["evidence_family"] not in allowed_families:
            normalized_cue["evidence_family"] = "care_context"
        if not normalized_cue["surface_form"]:
            normalized_cue["surface_form"] = "검토 필요 단서"
        selected.append(normalized_cue)
    withheld: list[dict[str, str]] = []
    allowed_withheld_reasons = {
        "answer_evidence", "redundant_support", "answer_revealing",
        "task_irrelevant", "unsupported_precision", "disclosure_budget",
        "unspecified_review",
    }
    for cue in withheld_raw:
        if isinstance(cue, str) and cue.strip():
            withheld.append(
                {
                    "source_id": cue.strip(),
                    "reason": "unspecified_review",
                }
            )
            continue
        if not isinstance(cue, dict):
            invalid_selected_cue = True
            continue
        source_id = str(cue.get("source_id") or "").strip()
        reason = str(cue.get("reason") or "").strip()
        if not source_id or reason not in allowed_withheld_reasons:
            invalid_selected_cue = True
        withheld.append(
            {
                "source_id": source_id or "unresolved_source",
                "reason": reason if reason in allowed_withheld_reasons else "task_irrelevant",
            }
        )

    allowed_qc_flags = {
        "diagnosis_leak", "overdetermined_diagnosis", "target_already_resolved",
        "redundant_confirmatory_evidence", "rag_to_stem_copy", "unsupported_precision",
        "task_evidence_mismatch", "malicious_red_herring", "missing_disclosure_plan",
        "disclosure_budget_exceeded", "patient_characteristic_bias",
    }
    raw_qc_flags = [str(flag).strip() for flag in raw.get("qc_flags") or [] if str(flag).strip()]
    qc_flags = [flag for flag in raw_qc_flags if flag in allowed_qc_flags]
    if any(flag not in allowed_qc_flags for flag in raw_qc_flags):
        qc_flags.append("task_evidence_mismatch")
    if missing and "missing_disclosure_plan" not in qc_flags:
        qc_flags.append("missing_disclosure_plan")
    if raw and not selected:
        qc_flags.append("task_evidence_mismatch")
    if invalid_selected_cue:
        qc_flags.append("task_evidence_mismatch")
    if len(selected_raw) > 6:
        qc_flags.append("disclosure_budget_exceeded")
    informative_count = sum(1 for cue in selected if cue["role"] != "neutral_context")
    strong_count = sum(1 for cue in selected if cue["strength"] in {"strong", "confirmatory"})
    confirmatory_count = sum(1 for cue in selected if cue["strength"] == "confirmatory")
    neutral_count = sum(1 for cue in selected if cue["role"] == "neutral_context")
    competing_count = sum(1 for cue in selected if cue["role"] == "authentic_competing_cue")
    if informative_count > 4 or strong_count > 1 or neutral_count > 1 or competing_count > 1:
        qc_flags.append("disclosure_budget_exceeded")
    if strong_count > 1 or confirmatory_count > 1:
        qc_flags.append("overdetermined_diagnosis")
    target_already_resolved = normalize_bool(raw.get("target_already_resolved"), False)
    if target_already_resolved:
        qc_flags.append("target_already_resolved")
    status = str(raw.get("status") or ("blocked" if qc_flags else "pass")).strip()
    if status not in {"pass", "blocked"}:
        status = "blocked"
        qc_flags.append("task_evidence_mismatch")
    if status == "blocked" and not qc_flags:
        qc_flags.append("task_evidence_mismatch")
    if qc_flags:
        status = "blocked"
    retrieved_raw = raw.get("retrieved_axis_ids") or []
    if isinstance(retrieved_raw, (str, bytes)):
        retrieved_raw = [retrieved_raw]
    retrieved_raw = [
        *retrieved_raw,
        *(cue["source_id"] for cue in selected if cue.get("source_id")),
    ]
    return {
        "status": status,
        "assessment_task": str(raw.get("assessment_task") or assessment_task or "diagnosis").strip(),
        "latent_diagnosis_required": normalize_bool(raw.get("latent_diagnosis_required"), False),
        "retrieved_axis_ids": sorted({str(value).strip() for value in retrieved_raw if str(value).strip()}),
        "selected_for_stem": selected,
        "withheld": withheld,
        "informative_cue_count": min(4, informative_count),
        "strong_or_confirmatory_count": min(1, strong_count),
        "confirmatory_count": min(1, confirmatory_count),
        "neutral_context_count": min(1, neutral_count),
        "authentic_competing_cue_count": min(1, competing_count),
        "target_already_resolved": target_already_resolved,
        "qc_flags": sorted(set(qc_flags)),
    }


def normalize_structured_choice_explanations(
    value: Any,
    *,
    answer_num: int,
    fallback: dict[str, str],
) -> dict[str, dict[str, str]]:
    raw = value if isinstance(value, dict) else {}
    out: dict[str, dict[str, str]] = {}
    for i in range(1, 6):
        key = str(i)
        entry = raw.get(key) or raw.get(i) or {}
        if isinstance(entry, dict):
            model_verdict = str(entry.get("verdict") or entry.get("declared_verdict") or "").strip()
            rationale = str(entry.get("rationale") or entry.get("correct_reason") or entry.get("explanation") or "").strip()
            misconception = str(entry.get("misconception") or "").strip()
            misconception_id = str(entry.get("misconception_id") or "").strip()
            why_attractive = str(entry.get("why_attractive") or "").strip()
            source_id = str(entry.get("source_id") or "").strip()
            provenance = str(entry.get("provenance") or "").strip()
            concept_level = str(entry.get("concept_level") or "").strip()
        else:
            model_verdict = ""
            rationale = str(entry or "").strip()
            misconception = ""
            misconception_id = ""
            why_attractive = ""
            source_id = ""
            provenance = ""
            concept_level = ""
        if not rationale:
            rationale = fallback.get(key, "")
        out[key] = {
            "verdict": "정답" if i == answer_num else "오답",
            # Preserve the model's declaration before deterministic
            # normalization so answer-key/explanation mismatch remains
            # detectable instead of being silently overwritten.
            "model_verdict": model_verdict,
            "rationale": rationale,
            "misconception": "" if i == answer_num else misconception,
            "misconception_id": "" if i == answer_num else misconception_id,
            "why_attractive": "" if i == answer_num else why_attractive,
            "source_id": "" if i == answer_num else source_id,
            "provenance": "" if i == answer_num else provenance,
            "concept_level": concept_level,
        }
    return out


def normalize_bool(value: Any, default: bool = False) -> bool:
    if isinstance(value, bool):
        return value
    if value is None:
        return default
    return str(value).strip().lower() in {"1", "true", "yes", "on"}


def normalize_data_table(value: Any) -> dict[str, Any] | None:
    if not value:
        return None
    if isinstance(value, dict):
        columns_raw = value.get("columns") or []
        rows_raw = value.get("rows") or []
        columns = [str(item).strip() for item in columns_raw if str(item).strip()] if isinstance(columns_raw, list) else []
        rows: list[list[str]] = []
        if isinstance(rows_raw, list):
            for row in rows_raw:
                if isinstance(row, list):
                    cells = [str(cell).strip() for cell in row]
                elif isinstance(row, dict):
                    cells = [str(row.get(column) or "").strip() for column in columns]
                else:
                    cells = [str(row).strip()]
                if any(cells):
                    rows.append(cells)
        if columns and rows:
            return {
                "title": str(value.get("title") or "").strip(),
                "columns": columns,
                "rows": rows[:12],
            }
    if isinstance(value, list):
        rows = [[str(cell).strip() for cell in row] for row in value if isinstance(row, list)]
        rows = [row for row in rows if any(row)]
        if len(rows) >= 2:
            return {
                "title": "",
                "columns": rows[0],
                "rows": rows[1:13],
            }
    return None


def normalize_question(
    item: dict[str, Any],
    *,
    idx: int,
    source_name: str,
    subject: str,
    unit: str,
    reveal_specialty: bool = False,
    reasoning_hops: int = 2,
    item_type: str = "A",
) -> dict[str, Any]:
    problem = str(item.get("problem") or item.get("stem") or item.get("question") or "").strip()
    options_raw = item.get("options") or item.get("choices") or []
    if isinstance(options_raw, dict):
        options = [str(options_raw.get(str(i)) or options_raw.get(i) or "").strip() for i in range(1, 6)]
    else:
        options = [str(opt).strip() for opt in list(options_raw)]
    options = [opt for opt in options if opt]
    while len(options) < 5:
        options.append(f"검수 필요 보기 {len(options) + 1}")
    options = options[:5]

    answer = item.get("answer")
    invalid_answer = False
    try:
        answer_num = int(answer)
    except Exception:
        invalid_answer = True
        answer_num = 1
    if answer_num < 1 or answer_num > 5:
        invalid_answer = True
        answer_num = 1

    pma_solution = item.get("pma_solution") if isinstance(item.get("pma_solution"), dict) else {}
    pma_solution = {
        "reasoning_summary": str(pma_solution.get("reasoning_summary") or "").strip(),
        "correct_reason": str(pma_solution.get("correct_reason") or "").strip(),
        "choice_explanations": normalize_choice_explanations(pma_solution.get("choice_explanations")),
        "high_yield_point": str(pma_solution.get("high_yield_point") or "").strip(),
        "trap": str(pma_solution.get("trap") or "").strip(),
        "source_anchor": str(pma_solution.get("source_anchor") or "").strip(),
    }
    cognitive_model = normalize_cognitive_model(item.get("cognitive_model"))
    item_question_type = str(item.get("question_type") or "clinical_case").strip()
    disclosure_plan = normalize_evidence_disclosure_plan(
        item.get("evidence_disclosure_plan"),
        assessment_task=item_question_type,
    )
    structured_choice_explanations = normalize_structured_choice_explanations(
        item.get("choice_explanations"),
        answer_num=answer_num,
        fallback=pma_solution["choice_explanations"],
    )
    try:
        normalized_hops = max(1, min(3, int(item.get("reasoning_hops") or reasoning_hops)))
    except (TypeError, ValueError):
        normalized_hops = max(1, min(3, int(reasoning_hops)))
    evidence_refs = item.get("evidence_refs")
    if not isinstance(evidence_refs, list):
        evidence_refs = []
    normalized_refs: list[dict[str, str]] = []
    allowed_source_types = {"lecture", "guideline", "journal", "textbook", "database", "other"}
    for ref in evidence_refs:
        if not isinstance(ref, dict):
            continue
        source_type = str(ref.get("source_type") or "other").strip()
        if source_type not in allowed_source_types:
            source_type = "other"
        normalized_refs.append(
            {
                "source_id": str(ref.get("source_id") or "").strip(),
                "source": str(ref.get("source") or "").strip(),
                "locator": str(ref.get("locator") or "").strip(),
                "chapter": str(ref.get("chapter") or "").strip(),
                "printed_page": str(ref.get("printed_page") or ref.get("page") or "").strip(),
                "basis": str(ref.get("basis") or "").strip(),
                "source_type": source_type,
                "retrieval_method": str(ref.get("retrieval_method") or "").strip(),
                "entailment_status": str(ref.get("entailment_status") or "").strip(),
            }
        )
    evidence_tier = str(item.get("evidence_tier") or "lecture_only").strip()
    if evidence_tier not in {
        "lecture_only",
        "lecture_plus_guideline",
        "lecture_plus_journal",
        "lecture_plus_textbook",
        "lecture_plus_database",
        "lecture_plus_mixed",
        "needs_external_review",
    }:
        evidence_tier = "needs_external_review"

    explanation = str(item.get("explanation") or "").strip()
    if not explanation:
        explanation = (
            f"풀이: {pma_solution['reasoning_summary']}\n"
            f"정답: {answer_num}번. {pma_solution['correct_reason']}\n"
            f"출제 포인트: {pma_solution['high_yield_point']}"
        ).strip()

    review_reasons: list[str] = []
    if not problem:
        review_reasons.append("empty_problem")
    if invalid_answer:
        review_reasons.append("invalid_answer")
    if len([opt for opt in options if not opt.startswith("검수 필요 보기")]) < 5:
        review_reasons.append("choice_count_lt_5")
    if not pma_solution["correct_reason"]:
        review_reasons.append("missing_correct_reason")
    if not pma_solution["source_anchor"]:
        review_reasons.append("missing_source_anchor")
    if bool(item.get("needs_review")):
        review_reasons.append("model_marked_needs_review")
    if evidence_tier == "needs_external_review":
        review_reasons.append("needs_external_evidence_review")
    if disclosure_plan["status"] != "pass":
        review_reasons.append("evidence_disclosure_plan_blocked")
    data_table = normalize_data_table(item.get("data_table") or item.get("table"))
    if data_table and bool(item.get("needs_review")):
        review_reasons.append("table_value_review_needed")
    review_reasons.append("automated_generation_requires_human_review")

    record = {
        "question_id": f"LECTURE_{slugify(Path(source_name).stem)}_Q{idx:03d}",
        "source_name": source_name,
        "source_type": "lecture_material",
        "subject": str(item.get("subject") or subject or "General").strip() or "General",
        "unit": str(item.get("unit") or unit or "미분류").strip() or "미분류",
        "problem": problem,
        "options": options,
        "answer": answer_num,
        "explanation": explanation,
        "pma_solution": pma_solution,
        "evidence_refs": normalized_refs,
        "evidence_tier": evidence_tier,
        "reference_notes": item.get("reference_notes", []),
        "difficulty": str(item.get("difficulty") or "").strip(),
        "question_type": item_question_type,
        "cognitive_level": str(item.get("cognitive_level") or "L3 Clinical Reasoning").strip(),
        "cognitive_model": cognitive_model,
        "evidence_disclosure_plan": disclosure_plan,
        "choice_explanations": structured_choice_explanations,
        "item_type": str(item.get("item_type") or item_type or "A").strip() or "A",
        "reveal_specialty": normalize_bool(item.get("reveal_specialty"), reveal_specialty),
        "reasoning_hops": normalized_hops,
        "self_check": item.get("self_check") if isinstance(item.get("self_check"), dict) else {},
        "review_status": "draft",
        "needs_review": True,
        "review_reasons": sorted(set(review_reasons)),
        "gen_ready": False,
        "block_reason": str(item.get("block_reason") or "").strip() or None,
        "generation_mode": "lecture_high_yield_pma_style",
        "pipeline_version": PIPELINE_VERSION,
    }
    if data_table:
        record["data_table"] = data_table
    return apply_generation_quality_gate(record)


def normalize_questions(
    payload: Any,
    *,
    source_name: str,
    subject: str,
    unit: str,
    reveal_specialty: bool = False,
    reasoning_hops: int = 2,
    item_type: str = "A",
) -> list[dict[str, Any]]:
    if isinstance(payload, dict):
        payload = payload.get("questions") or payload.get("items") or [payload]
    if not isinstance(payload, list):
        raise ValueError("JSON payload must be a list or object containing questions/items")
    out: list[dict[str, Any]] = []
    for idx, item in enumerate(payload, 1):
        if isinstance(item, dict):
            out.append(
                normalize_question(
                    item,
                    idx=idx,
                    source_name=source_name,
                    subject=subject,
                    unit=unit,
                    reveal_specialty=reveal_specialty,
                    reasoning_hops=reasoning_hops,
                    item_type=item_type,
                )
            )
    return out


def generate_openai(prompt: str, *, model: str, temperature: float) -> str:
    from openai import OpenAI

    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        raise RuntimeError("OPENAI_API_KEY is not set")
    client = OpenAI(api_key=api_key)
    response = client.chat.completions.create(
        model=model,
        messages=[
            {"role": "system", "content": f"{GENERATION_SYSTEM_PROMPT}\n\nReturn only valid JSON. Do not include markdown fences."},
            {"role": "user", "content": prompt},
        ],
        temperature=temperature,
    )
    return response.choices[0].message.content or ""


def generate_gemini(prompt: str, *, model: str, temperature: float) -> str:
    import google.generativeai as genai

    api_key = os.getenv("GEMINI_API_KEY")
    if not api_key:
        raise RuntimeError("GEMINI_API_KEY is not set")
    genai.configure(api_key=api_key)
    gemini_model = genai.GenerativeModel(model)
    response = gemini_model.generate_content(
        prompt,
        generation_config={"temperature": temperature, "top_p": 1.0},
    )
    return response.text or ""


def resolve_provider(provider: str) -> str:
    if provider != "auto":
        return provider
    if os.getenv("OPENAI_API_KEY"):
        return "openai"
    if os.getenv("GEMINI_API_KEY"):
        return "gemini"
    return "prompt-only"


def print_sample(records: list[dict[str, Any]], limit: int) -> None:
    if limit <= 0:
        return
    print(f"\nSAMPLE {min(limit, len(records))}/{len(records)}")
    for record in records[:limit]:
        preview = record.get("problem", "")[:80].replace("\n", " ")
        if len(record.get("problem", "")) > 80:
            preview += "..."
        print(
            f"- {record.get('question_id')}: answer={record.get('answer')} "
            f"needs_review={record.get('needs_review')} stem={preview}"
        )


def process_file(path: Path, args: argparse.Namespace) -> Path:
    text = extract_text(path)
    if not text:
        raise RuntimeError(f"no text extracted from {path}")

    args.extracted_dir.mkdir(parents=True, exist_ok=True)
    args.output_dir.mkdir(parents=True, exist_ok=True)

    slug = slugify(path.stem)
    extracted_path = args.extracted_dir / f"{slug}.txt"
    prompt_path = args.output_dir / f"{slug}.prompt.txt"
    output_path = args.output_dir / f"{slug}.questions.json"
    raw_response_path = args.output_dir / f"{slug}.model_response.txt"

    extracted_path.write_text(text, encoding="utf-8")
    evidence_paths = collect_evidence_paths(args)
    evidence_context = build_evidence_context(
        evidence_paths,
        max_chars_per_file=args.evidence_max_chars_per_file,
        max_total_chars=args.evidence_max_total_chars,
    )
    prompt = build_generation_prompt(
        text,
        source_name=path.name,
        subject=args.subject,
        unit=args.unit,
        num_questions=args.num_questions,
        difficulty=args.difficulty,
        max_chars=args.max_chars,
        evidence_context=evidence_context,
        reveal_specialty=getattr(args, "reveal_specialty", False),
        reasoning_hops=getattr(args, "reasoning_hops", 2),
        item_type=getattr(args, "item_type", "A"),
        exam_profile=getattr(args, "exam_profile", "kmle_summative"),
        assessment_task=getattr(args, "assessment_task", "diagnosis"),
        question_type=getattr(args, "question_type", "clinical_case"),
    )
    grounding = build_generation_grounding(
        args.unit,
        disease_concept_id=getattr(args, "disease_concept_id", ""),
        review_policy=getattr(args, "ontology_review_policy", "faculty_draft"),
        retrieval_intents=[getattr(args, "assessment_task", "diagnosis")],
    )
    prompt = append_grounding_context(prompt, grounding)
    prompt_path.write_text(prompt, encoding="utf-8")

    # Fail closed before any model call.  This prevents a plausible-looking
    # question from being generated when matching, four homogeneous ontology
    # distractors, or Harrison/approved evidence are missing.
    standalone_block_reasons = list(
        grounding.get("block_reasons")
        or grounding.get("draft_generation_block_reasons")
        or []
    )
    if grounding.get("blocked") or standalone_block_reasons:
        packet = {
            "source_name": path.name,
            "provider": resolve_provider(args.provider),
            "status": "grounding_blocked",
            "block_reasons": standalone_block_reasons,
            "missing": grounding.get("missing") or [],
            "match": grounding.get("match") or {},
            "prompt_path": str(prompt_path),
            "extracted_text_path": str(extracted_path),
            "created_at": datetime.now(timezone.utc).isoformat(),
            "pipeline_version": PIPELINE_VERSION,
            "needs_review": True,
            "gen_ready": False,
        }
        output_path.write_text(json.dumps(packet, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"[grounding-blocked] {path.name}: {packet['block_reasons']}", file=sys.stderr)
        if getattr(args, "validate", False):
            validate_output(output_path)
        return output_path

    provider = resolve_provider(args.provider)
    if provider == "prompt-only":
        packet = {
            "source_name": path.name,
            "provider": provider,
            "status": "prompt_ready",
            "prompt_path": str(prompt_path),
            "extracted_text_path": str(extracted_path),
            "created_at": datetime.now(timezone.utc).isoformat(),
            "pipeline_version": PIPELINE_VERSION,
            "needs_review": True,
            "gen_ready": False,
            "grounding": {
                "match": grounding.get("match") or {},
                "missing": grounding.get("missing") or [],
                "disease_concept_id": ((grounding.get("pack") or {}).get("disease_concept_id")),
            },
        }
        output_path.write_text(json.dumps(packet, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"[prompt-only] {path.name} -> {prompt_path}", file=sys.stderr)
        if getattr(args, "validate", False):
            validate_output(output_path)
        return output_path

    model = args.openai_model if provider == "openai" else args.gemini_model
    if provider == "openai":
        raw_response = generate_openai(prompt, model=model, temperature=args.temperature)
    elif provider == "gemini":
        raw_response = generate_gemini(prompt, model=model, temperature=args.temperature)
    else:
        raise RuntimeError(f"Unsupported provider: {provider}")

    raw_response_path.write_text(raw_response, encoding="utf-8")
    payload = extract_json_payload(raw_response)
    records = normalize_questions(
        payload,
        source_name=path.name,
        subject=args.subject,
        unit=args.unit,
        reveal_specialty=getattr(args, "reveal_specialty", False),
        reasoning_hops=getattr(args, "reasoning_hops", 2),
        item_type=getattr(args, "item_type", "A"),
    )
    records = [apply_grounding_trace(record, grounding) for record in records]
    output_path.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
    print_sample(records, args.sample)
    print(f"[generated] {path.name} -> {output_path}", file=sys.stderr)
    if getattr(args, "validate", False):
        validate_output(output_path)
    return output_path


def validate_output(output_path: Path) -> None:
    """Validate output_path JSON against lecture_question.schema.json.

    - Skips prompt_ready packets (dict with status == "prompt_ready").
    - Raises jsonschema.ValidationError on schema failure.
    """
    try:
        import jsonschema
    except ImportError:  # pragma: no cover
        print("[validate] jsonschema not installed — skipping validation.", file=sys.stderr)
        return

    schema_path = Path(__file__).parent.parent / "schemas" / "lecture_question.schema.json"
    if not schema_path.exists():
        print(f"[validate] schema not found: {schema_path}", file=sys.stderr)
        return

    schema = json.loads(schema_path.read_text(encoding="utf-8"))
    data = json.loads(output_path.read_text(encoding="utf-8"))

    if isinstance(data, dict) and data.get("status") in {"prompt_ready", "grounding_blocked"}:
        print(
            f"[validate] skip — {data.get('status')} packet (no questions generated yet): {output_path.name}",
            file=sys.stderr,
        )
        return

    jsonschema.validate(instance=data, schema=schema)
    count = len(data) if isinstance(data, list) else 1
    print(f"[validate] OK — {count} records passed schema: {output_path.name}", file=sys.stderr)


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Generate PMA-style MCQ draft JSON from lecture files.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    parser.add_argument("inputs", nargs="+", help="Lecture file paths: pdf/docx/pptx/txt/md/hwp")
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--extracted-dir", type=Path, default=DEFAULT_EXTRACTED_DIR)
    parser.add_argument("--provider", choices=["auto", "prompt-only", "openai", "gemini"], default="auto")
    parser.add_argument("--openai-model", default=os.getenv("OPENAI_MODEL", "gpt-4o-mini"))
    parser.add_argument("--gemini-model", default=os.getenv("GEMINI_MODEL", "gemini-2.5-flash"))
    parser.add_argument("--num-questions", type=int, default=5)
    parser.add_argument("--subject", default="General")
    parser.add_argument("--unit", default="미분류")
    parser.add_argument("--difficulty", default="보통")
    parser.add_argument("--disease-concept-id", default="", help="Exact reviewed concept_registry ID.")
    parser.add_argument(
        "--ontology-review-policy",
        choices=["faculty_draft", "student_approved"],
        default="faculty_draft",
    )
    parser.add_argument(
        "--assessment-task",
        choices=["classification", "diagnosis", "treatment", "mechanism"],
        default="diagnosis",
    )
    parser.add_argument("--question-type", choices=["clinical_case", "concept"], default="clinical_case")
    parser.add_argument("--item-type", default="A")
    parser.add_argument("--exam-profile", default="kmle_summative")
    parser.add_argument("--reasoning-hops", type=int, choices=[1, 2, 3], default=2)
    parser.add_argument("--reveal-specialty", action="store_true", default=False)
    parser.add_argument("--max-chars", type=int, default=30000)
    parser.add_argument("--temperature", type=float, default=0.0)
    parser.add_argument("--sample", type=int, default=5)
    parser.add_argument(
        "--evidence",
        action="append",
        default=[],
        help="Approved local evidence file to use as supplemental context. Repeatable.",
    )
    parser.add_argument(
        "--evidence-dir",
        action="append",
        default=[],
        help="Directory containing approved local evidence files. Repeatable.",
    )
    parser.add_argument("--evidence-max-chars-per-file", type=int, default=8000)
    parser.add_argument("--evidence-max-total-chars", type=int, default=20000)
    parser.add_argument(
        "--validate",
        action="store_true",
        default=False,
        help=(
            "After generation, validate the output JSON against "
            "schemas/lecture_question.schema.json. "
            "prompt_ready packets are skipped automatically."
        ),
    )
    return parser


def main() -> None:
    args = build_arg_parser().parse_args()
    check_gitignore()
    for raw in args.inputs:
        path = Path(raw).expanduser()
        if not path.exists():
            raise SystemExit(f"Input not found: {path}")
        process_file(path, args)


if __name__ == "__main__":
    main()
