#!/usr/bin/env python3
"""생성 문항 JSON → 해설집 DOCX (교시별).

시험지 본문과 정답표에는 해설이 들어가지 않는다(배포 안전). 검수자·교수 검토용으로
문항별 전체 근거를 한 권으로 묶는다.

구성(문항당):
  [No n] 정답 + 분과·평가축·인지수준·추론단계
  문두(요약) → 검사결과 표 → 선지 5개(정답 표시) → 해설 → 선지별 분석 → 근거 → 온톨로지 개념
입력: --items data_private/professor_items/generated/set_<k>.json
출력: exports/<제목>_해설집.docx
"""
import argparse
import json
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

try:
    from scripts.lab_box_render import add_lab_box
except ImportError:  # 직접 실행(python3 scripts/…) 시
    from lab_box_render import add_lab_box

BASE = Path("data_private/professor_items")
IMG_DIR = BASE / "images"
# 웹 배포용으로 무손실 PNG 변환해 둔 사본(원본 BMP는 장당 수십 MB라 문서에 부적합)
PNG_DIR = Path("data_private/course_exams/media/SYNTH_GENERATED")


def resolve_image(name: str):
    """최적화 PNG 우선, 없으면 원본. 문서 용량을 수백 MB에서 수십 MB로 낮춘다."""
    if not name:
        return None
    png = PNG_DIR / (Path(name).stem + ".png")
    if png.exists():
        return png
    for cand in (PNG_DIR / name, IMG_DIR / name):
        if cand.exists():
            return cand
    return None
OUT = BASE / "exports"


CIRC = {"1": "①", "2": "②", "3": "③", "4": "④", "5": "⑤"}
BOOK_TITLES = {
    "harrison_22e": "Harrison 22e", "sabiston_21e": "Sabiston 21e",
    "nelson_21e": "Nelson 21e", "williams_ob_25e": "Williams Obstetrics 25e",
    "berek_novak_16e": "Berek & Novak 16e", "speroff_9e": "Speroff",
    "kr_psychiatry_3e": "신경정신의학 3판",
}
TEAL = RGBColor(0x0E, 0x7C, 0x7B)
GREY = RGBColor(0x64, 0x74, 0x8B)
RED = RGBColor(0xA1, 0x3F, 0x31)

VERDICT_BADGE = {
    "fully": ("✓ 근거 페이지 검증", TEAL),
    "rerouted": ("◇ 근거 재지정(과별 교과서) — 검증 전", GREY),
    "partially": ("△ 핵심 확인 — 일부 주장은 인용 장 밖(감별질환 등)", RGBColor(0xB4, 0x53, 0x09)),
    "not": ("✗ 인용 장에서 미확인 — 근거 재지정·의학 검수 필요", RED),
}


def head(doc, text, size=10, color=TEAL):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = color
    return p


def body(doc, text, size=9.5, indent=0.4):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(indent)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(str(text))
    r.font.size = Pt(size)
    return p


def add_item(doc, it):
    ans = str(it.get("answer", ""))
    p = doc.add_paragraph()
    r = p.add_run(f"[No {it['no']}] 정답 {CIRC.get(ans, ans)}")
    r.bold = True
    r.font.size = Pt(12)
    meta = " · ".join(x for x in [
        it.get("subject", ""), it.get("axis", ""),
        it.get("cognitive_level", ""),
        f"추론 {it['reasoning_hops']}단계" if it.get("reasoning_hops") else "",
        "이미지 문항" if it.get("image") else "",
    ] if x)
    if meta:
        mr = p.add_run(f"    {meta}")
        mr.font.size = Pt(8.5)
        mr.font.color.rgb = GREY

    body(doc, str(it.get("stem", "")), size=9, indent=0.2)

    if it.get("lab_box"):
        # 종전엔 해설집에 검사수치가 아예 빠져 있었다. 설문 T3 반영해 2열 표로
        # 넣고, 파싱 실패 시 원문 그대로 1칸 박스 폴백(유실 없음).
        add_lab_box(doc, it["lab_box"], fallback_size=9)

    src_img = resolve_image(str(it.get("image") or ""))
    if src_img:
        ip = doc.add_paragraph()
        ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
        ip.paragraph_format.space_after = Pt(2)
        try:
            ip.add_run().add_picture(str(src_img), width=Cm(8.5))
        except Exception:
            ip.add_run(f"[이미지: {it.get('image')}]").font.size = Pt(8)

    for k in sorted(it.get("choices", {}), key=lambda x: int(x)):
        mark = " ◀ 정답" if k == ans else ""
        pp = doc.add_paragraph()
        pp.paragraph_format.left_indent = Cm(0.6)
        pp.paragraph_format.space_after = Pt(0)
        rr = pp.add_run(f"{CIRC.get(k, k)} {it['choices'][k]}{mark}")
        rr.font.size = Pt(9)
        if k == ans:
            rr.bold = True

    if it.get("explanation"):
        head(doc, "해설")
        body(doc, it["explanation"])

    ce = it.get("choice_explanations") or {}
    if ce:
        head(doc, "선지별 분석")
        for k in sorted(ce, key=lambda x: int(x)):
            row = ce[k] or {}
            txt = str(row.get("why_correct") or row.get("why_attractive") or "").strip()
            if not txt:
                continue
            pp = doc.add_paragraph()
            pp.paragraph_format.left_indent = Cm(0.6)
            pp.paragraph_format.space_after = Pt(1)
            lab = pp.add_run(f"{CIRC.get(k, k)} ")
            lab.bold = True
            lab.font.size = Pt(9)
            if k == ans:
                lab.font.color.rgb = TEAL
            rr = pp.add_run(txt)
            rr.font.size = Pt(9)

    src = it.get("harrison_sources") or []
    tsrc = it.get("textbook_sources") or []
    cid = it.get("disease_concept_id") or ""
    if src or tsrc or cid:
        bits = []
        for s in src:
            bits.append(f"Harrison 22e Ch.{s.get('chapter')} p.{s.get('printed_page')}"
                        f"{' (장 포인터, 주장 검증 전)' if s.get('entailment_status') == 'chapter_pointer' else ''}")
        for t in tsrc:
            bits.append(f"{BOOK_TITLES.get(str(t.get('book_id')), t.get('book_id'))} "
                        f"Ch.{t.get('chapter')} (장 포인터)")
        if cid:
            bits.append(f"온톨로지 개념: {cid}")
        head(doc, "근거")
        body(doc, " · ".join(bits))
        badge = VERDICT_BADGE.get(str(it.get("entailment_verdict") or ""))
        if badge:
            bp = doc.add_paragraph()
            bp.paragraph_format.left_indent = Cm(0.4)
            br = bp.add_run(badge[0])
            br.font.size = Pt(8.5)
            br.bold = True
            br.font.color.rgb = badge[1]
            for pt in (it.get("entailment_unsupported_points") or [])[:3]:
                pp2 = doc.add_paragraph()
                pp2.paragraph_format.left_indent = Cm(0.7)
                rr2 = pp2.add_run(f"· 확인 필요: {pt}")
                rr2.font.size = Pt(8)
                rr2.font.color.rgb = GREY

    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--items", required=True)
    ap.add_argument("--title", default="해설집")
    args = ap.parse_args()

    items = json.loads(Path(args.items).read_text(encoding="utf-8"))
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(1.8)
        s.left_margin = s.right_margin = Cm(1.8)
    style = doc.styles["Normal"]
    style.font.name = "맑은 고딕"
    style.font.size = Pt(9.5)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = t.add_run(f"{args.title} — 해설집")
    tr.bold = True
    tr.font.size = Pt(16)
    n = doc.add_paragraph()
    n.alignment = WD_ALIGN_PARAGRAPH.CENTER
    nr = n.add_run("검수·교수 검토용 · 학생 배포 금지 · 의학적 검수 전 AI 생성물")
    nr.font.size = Pt(9)
    nr.font.color.rgb = RGBColor(0xA1, 0x3F, 0x31)
    doc.add_paragraph()

    for it in items:
        add_item(doc, it)

    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / f"{args.title}_해설집.docx"
    doc.save(path)
    ce_full = sum(1 for it in items if len(it.get("choice_explanations") or {}) == 5)
    src_n = sum(1 for it in items if it.get("harrison_sources"))
    img_n = sum(1 for it in items if resolve_image(str(it.get("image") or "")))
    want_img = sum(1 for it in items if it.get("image"))
    print(f"[해설집] {path}  ({len(items)}문항 · 선지별해설 {ce_full} · 근거 {src_n} "
          f"· 이미지 {img_n}/{want_img} · {path.stat().st_size // 1024 // 1024} MB)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
