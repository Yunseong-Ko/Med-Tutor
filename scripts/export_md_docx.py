#!/usr/bin/env python3
"""마크다운 문서 → 교수 전달용 DOCX 변환 (경량 렌더러).

지원: #/##/### 제목, 불릿(-), 번호목록(1.), 표(|…|), **굵게** 인라인, 문단.
프로젝트 문서 톤(맑은 고딕, 틸 헤딩)을 기존 익스포터와 맞춘다.
사용: python3 scripts/export_md_docx.py --md docs/문서.md --out exports/문서.docx
"""
import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

TEAL = RGBColor(0x0E, 0x7C, 0x7B)
GREY = RGBColor(0x64, 0x74, 0x8B)
BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def add_runs(p, text, size=10):
    """**굵게** 인라인을 살려 런 분할."""
    pos = 0
    for m in BOLD_RE.finditer(text):
        if m.start() > pos:
            r = p.add_run(text[pos:m.start()])
            r.font.size = Pt(size)
        r = p.add_run(m.group(1))
        r.bold = True
        r.font.size = Pt(size)
        pos = m.end()
    if pos < len(text):
        r = p.add_run(text[pos:])
        r.font.size = Pt(size)


def render(md_path: Path, out_path: Path):
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2.0)
        s.left_margin = s.right_margin = Cm(2.0)
    doc.styles["Normal"].font.name = "맑은 고딕"
    doc.styles["Normal"].font.size = Pt(10)

    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln.strip():
            i += 1
            continue
        if ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            if rows:
                ncols = max(len(r) for r in rows)
                t = doc.add_table(rows=len(rows), cols=ncols)
                t.style = "Table Grid"
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                for ri, row in enumerate(rows):
                    for ci in range(ncols):
                        cell = t.cell(ri, ci)
                        cell.paragraphs[0].text = ""
                        add_runs(cell.paragraphs[0],
                                 row[ci] if ci < len(row) else "",
                                 size=8.5 if ri else 9)
                        if ri == 0:
                            for r in cell.paragraphs[0].runs:
                                r.bold = True
                doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue
        if ln.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            r = p.add_run(ln[4:])
            r.bold = True
            r.font.size = Pt(11)
            r.font.color.rgb = TEAL
        elif ln.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            r = p.add_run(ln[3:])
            r.bold = True
            r.font.size = Pt(13)
            r.font.color.rgb = TEAL
        elif ln.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(ln[2:])
            r.bold = True
            r.font.size = Pt(16)
        elif re.match(r"^\d+\.\s", ln.strip()):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, ln.strip())
        elif ln.strip().startswith("- "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, "· " + ln.strip()[2:])
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            add_runs(p, ln.strip())
        i += 1

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    print(f"[DOCX] {out_path}  ({out_path.stat().st_size // 1024} KB)")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--md", required=True)
    ap.add_argument("--out", required=True)
    args = ap.parse_args()
    render(Path(args.md), Path(args.out))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
