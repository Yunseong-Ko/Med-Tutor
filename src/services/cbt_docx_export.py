"""의학교육실 과정시험 HWP 양식 .docx 익스포터.

참조 양식(혈액 및 종양학 과정시험)을 1:1로 재현한다:

    @ {학년도} {과목} {차수} 과정시험/@      ← 시험지 헤더
    @ {출제교수}/@                            ← 출제자별 그룹
    N. {문두} {정답동그라미}                   ← 문두 끝 정답 마커
    <그림>                                    ← 이미지(있으면 삽입, 없으면 자리표시자)
    ① ... ⑤                                  ← 보기
    해설) {해설}                              ← 해설(우리 구조화 해설을 일관 출력)

한글(HWP)에서 .docx를 열어 .hwp로 저장하는 제출 준비용 산출물.
.hwp 직접 출력(hwp5)은 쓰기 지원이 약해 .docx 경유를 표준으로 한다.

사용:
    python3 -m src.services.cbt_docx_export <exam.json> [--out FILE]
        [--include-needs-review] [--group-by course_name|none]
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

CIRCLED = {"1": "①", "2": "②", "3": "③", "4": "④", "5": "⑤"}

# python-docx는 PyMuPDF가 추출한 일부 JPEG 헤더를 인식하지 못한다(UnrecognizedImageError).
# PIL로 표준 PNG로 재인코딩해 캐시한 뒤 임베드한다.
_IMG_CACHE = Path("data_private/course_exams/media_docx_cache")


def normalized_image(path: Path) -> Path | None:
    try:
        from PIL import Image
    except Exception:
        return path  # PIL 없으면 원본 시도
    try:
        _IMG_CACHE.mkdir(parents=True, exist_ok=True)
        out = _IMG_CACHE / (path.stem + ".png")
        if not out.exists():
            with Image.open(path) as im:
                im.convert("RGB").save(out, format="PNG")
        return out
    except Exception:
        return None


def circled(key: str) -> str:
    return CIRCLED.get(str(key).strip(), str(key))


def exam_header(exam: dict[str, Any]) -> str:
    parts = [exam.get("exam_title")]
    if not parts[0]:
        parts = [
            exam.get("exam_date", ""),
            exam.get("course_name", ""),
            exam.get("round_label", ""),
            exam.get("period_label", ""),
            "과정시험",
        ]
    return " ".join(p for p in parts if p).strip()


def group_key(question: dict[str, Any], mode: str) -> str:
    if mode == "none":
        return ""
    labels = question.get("labels") or {}
    # 출제교수 필드가 생기면 우선 사용, 없으면 과목명으로 그룹(양식의 '@ 교수님/@' 자리)
    return question.get("author") or labels.get("course_name") or "미분류"


def linked_image_path(question: dict[str, Any]) -> Path | None:
    media = question.get("media") or {}
    for ref in media.get("media_refs") or []:
        fp = ref.get("file_path")
        if fp and Path(fp).exists():
            return Path(fp)
    return None


# --- 모델 호환 접근자 (course_exam: stem/choices, studio: problem/options) ---
def _flat(s: Any) -> str:
    return str(s or "").replace("\n", " ").strip()


def q_stem(q: dict[str, Any]) -> str:
    from src.services.cbt_hwp_export import block_text
    return _flat(q.get("stem") or block_text(q.get("problem")))


def q_choice_rows(q: dict[str, Any]) -> list[tuple[str, str]]:
    """[(라벨①, 텍스트)] — choices(dict) 또는 options(list) 모두 처리."""
    from src.services.cbt_hwp_export import normalize_choices
    return normalize_choices(q)


def q_answer_mark(q: dict[str, Any]) -> str:
    from src.services.cbt_hwp_export import answer_label
    return answer_label(q)


def q_explanation(q: dict[str, Any]) -> str:
    from src.services.cbt_hwp_export import block_text
    if q.get("answer_rationale"):
        return _flat(q["answer_rationale"])
    ki = q.get("key_info") or {}
    if ki.get("core_explanation"):
        return _flat(ki["core_explanation"])
    for k in ("explanation", "pma_solution"):
        v = q.get(k)
        if v:
            return _flat(v if isinstance(v, str) else block_text(v))
    return ""


def is_essay(q: dict[str, Any]) -> bool:
    return (q.get("question_format") or "").lower() in {"essay", "subjective", "주관식"}


def points_label(q: dict[str, Any], exam: dict[str, Any]) -> str:
    pts = q.get("points")
    if pts is None:
        pts = exam.get("essay_default_points") if is_essay(q) else exam.get("objective_default_points")
    return f"({pts}점)" if pts is not None else ""


def evidence_note(q: dict[str, Any]) -> str:
    """출제근거 → '(출처: ...)'. 모드 A 변형은 부모 문항을 출처로 표기."""
    refs = q.get("evidence_refs") or []
    parts = []
    for r in refs:
        if isinstance(r, dict):
            parts.append(r.get("citation") or r.get("source") or r.get("page") or "")
        elif r:
            parts.append(str(r))
    parts = [p for p in parts if p]
    if not parts and q.get("parent_question_id"):
        parts.append(f"기출 변형: {q['parent_question_id']}")
    return f"(출처: {'; '.join(parts)})" if parts else ""


def gate_ok(q: dict[str, Any], gate: str) -> bool:
    if gate == "approved":
        return (q.get("review_status") or "").lower() == "approved"
    if gate == "all":
        return True
    return not q.get("needs_review")  # 기본: needs_review=false


def add_essay(doc: Document, q: dict[str, Any], exam: dict[str, Any], display_num: int | None = None) -> None:
    doc.add_paragraph("<-주관식->")
    num = q.get("question_number") or display_num
    stem = q_stem(q)
    pts = points_label(q, exam)
    p = doc.add_paragraph()
    r = p.add_run(f"{num}. {stem} {pts}".rstrip())
    r.font.size = Pt(10)

    media = q.get("media") or {}
    if media.get("has_image_or_data_reference") and media.get("media_refs"):
        img = linked_image_path(q)
        norm = normalized_image(img) if img is not None else None
        if norm is not None:
            try:
                doc.add_picture(str(norm), width=Inches(2.6))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            except Exception:
                doc.add_paragraph("<그림>")
        else:
            doc.add_paragraph("<그림>")

    ans = q.get("answer")
    ans_txt = ", ".join(ans) if isinstance(ans, list) else (ans or "")
    note = evidence_note(q)
    ap = doc.add_paragraph()
    ar = ap.add_run(f"정답: {ans_txt}   {note}".rstrip())
    ar.font.size = Pt(10)
    doc.add_paragraph("</-주관식->")
    doc.add_paragraph("")


def add_question(doc: Document, q: dict[str, Any], include_explanation: bool, display_num: int | None = None) -> None:
    num = q.get("question_number") or display_num
    stem = q_stem(q)
    ans = q_answer_mark(q)

    p = doc.add_paragraph()
    run = p.add_run(f"{num}. {stem} ")
    run.font.size = Pt(10)
    amark = p.add_run(ans)
    amark.bold = True
    amark.font.size = Pt(10)

    # 이미지: 연결된 자산이 있으면 삽입, 없으면 <그림> 자리표시자
    media = q.get("media") or {}
    if media.get("has_image_or_data_reference") and media.get("media_refs"):
        img = linked_image_path(q)
        norm = normalized_image(img) if img is not None else None
        if norm is not None:
            try:
                doc.add_picture(str(norm), width=Inches(2.6))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            except Exception:
                doc.add_paragraph("<그림>")
        else:
            doc.add_paragraph("<그림>")

    # 보기 (choices/options 모두 처리)
    for label, text in q_choice_rows(q):
        cp = doc.add_paragraph()
        cp.paragraph_format.left_indent = Pt(12)
        cr = cp.add_run(f"{label} {text}")
        cr.font.size = Pt(10)

    # 해설 (구조화 해설 → 일관 출력)
    if include_explanation:
        expl = q_explanation(q)
        if expl:
            ep = doc.add_paragraph()
            er = ep.add_run("해설) ")
            er.bold = True
            er.font.size = Pt(9)
            et = ep.add_run(expl.replace("\n", " ").strip())
            et.font.size = Pt(9)
            et.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    note = evidence_note(q)
    if note:
        np_ = doc.add_paragraph()
        nr = np_.add_run(note)
        nr.font.size = Pt(8)
        nr.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    doc.add_paragraph("")


# 기본 폰트: 이 macOS 환경에는 "맑은 고딕"/"함초롬"이 설치돼 있지 않아, 지정해도
# 뷰어(LibreOffice/한글/Word)가 제각각 다른 대체 폰트로 바꿔치기한다(숫자만 이상한
# 필기체로 보이는 등). "Apple SD Gothic Neo"는 이 Mac의 시스템 폰트라 어느 뷰어에서
# 열어도 정확히 일치해 렌더된다. Windows/HWP 환경에 배포할 때는 "맑은 고딕"으로
# 바꿔도 되지만, 그 폰트가 실제로 설치된 환경인지 먼저 확인할 것.
DEFAULT_FONT = "Apple SD Gothic Neo"


def set_font(font_or_run, name: str = DEFAULT_FONT) -> None:
    """ascii/hAnsi/eastAsia/cs 4개 슬롯을 모두 같은 폰트로 지정.

    python-docx의 Font.name은 ascii/hAnsi(라틴 문자)만 설정하고 eastAsia(한글)는
    건드리지 않는다. eastAsia를 비워두면 뷰어(HWP/Word/LibreOffice)가 스크립트별로
    서로 다른 대체 폰트를 골라 한글은 정상, 숫자·영문은 엉뚱한 폰트로 보이는 문제가
    생긴다. 네 슬롯을 전부 명시해 뷰어 간 렌더링 차이를 없앤다.
    """
    font_or_run.name = name
    rpr = font_or_run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)


def _render_document(exam, objective, essay, include_explanation, group_by) -> Document:
    doc = Document()
    style = doc.styles["Normal"]
    set_font(style.font, DEFAULT_FONT)
    style.font.size = Pt(10)

    htitle = doc.add_paragraph()
    hr = htitle.add_run(f"@ {exam_header(exam)}/@")
    hr.bold = True
    hr.font.size = Pt(13)
    summary = doc.add_paragraph()
    sr = summary.add_run(f"객관식 {len(objective)}문항 · 주관식 {len(essay)}문항")
    sr.font.size = Pt(9)
    sr.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    doc.add_paragraph("")

    def render_grouped(items, renderer):
        current = None
        for idx, q in enumerate(items, start=1):
            gk = group_key(q, group_by)
            if gk and gk != current:
                current = gk
                gp = doc.add_paragraph()
                gr = gp.add_run(f"@ {gk}/@")
                gr.bold = True
                gr.font.size = Pt(11)
            renderer(q, idx)

    render_grouped(objective, lambda q, n: add_question(doc, q, include_explanation, n))
    if essay:
        sp = doc.add_paragraph()
        sr2 = sp.add_run("@ 주관식/@")
        sr2.bold = True
        sr2.font.size = Pt(12)
        render_grouped(essay, lambda q, n: add_essay(doc, q, exam, n))
    return doc


def build_studio_cbt_docx_export(
    set_id: str,
    *,
    include_unapproved: bool = False,
    include_explanation: bool = True,
    group_by: str = "course_name",
) -> dict[str, Any]:
    """faculty-studio question-set → 의학교육실 양식 .docx (기존 cbt-hwp 라우트와 동일 모델)."""
    from datetime import datetime, timezone
    from src.services.cbt_hwp_export import safe_slug
    from src.services.lecture_studio import (
        EXPORT_SET_DIR,
        ensure_studio_dirs,
        load_question_set,
        write_json,
    )

    ensure_studio_dirs()
    packet = load_question_set(set_id, include_summary=True)
    metadata = packet.get("metadata") if isinstance(packet.get("metadata"), dict) else {}
    questions = packet.get("questions") if isinstance(packet.get("questions"), list) else []
    selected = [
        q for q in questions
        if include_unapproved or q.get("review_status") == "approved"
    ]
    if not selected:
        raise ValueError("승인된 문항이 없습니다. 문항 검토 큐에서 승인 후 의학교육실 HWP 양식을 생성하세요.")
    selected.sort(key=lambda q: q.get("question_number", 0))
    objective = [q for q in selected if not is_essay(q)]
    essay = [q for q in selected if is_essay(q)]

    exam = {
        "exam_title": metadata.get("exam_title"),
        "course_name": metadata.get("subject"),
        "period_label": metadata.get("unit"),
    }
    doc = _render_document(exam, objective, essay, include_explanation, group_by)

    slug = f"{safe_slug(set_id)}_cbt_docx"
    out = EXPORT_SET_DIR / f"{slug}.docx"
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))

    artifact = {
        "export_id": slug,
        "export_type": "cbt_docx_template",
        "set_id": set_id,
        "subject": metadata.get("subject"),
        "unit": metadata.get("unit"),
        "objective_count": len(objective),
        "essay_count": len(essay),
        "include_unapproved": include_unapproved,
        "file_path": str(out),
        "download_url": f"/api/exports/{out.name}",
        "hwp_workflow": "한글에서 .docx 열기 → 필요 시 편집 → 다른 이름으로 저장(.hwp)",
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    write_json(EXPORT_SET_DIR / f"{slug}.manifest.json", artifact)
    return artifact


def build_cbt_docx_export(
    exam_json_path: str | Path,
    out_path: str | Path | None = None,
    *,
    include_needs_review: bool = False,
    include_explanation: bool = True,
    group_by: str = "course_name",
    gate: str | None = None,
) -> dict[str, Any]:
    data = json.loads(Path(exam_json_path).read_text(encoding="utf-8"))
    exam = data.get("exam") or {}
    questions = data.get("questions") or []

    # 게이트: gate 우선, 없으면 include_needs_review로 결정(하위호환)
    if gate is None:
        gate = "all" if include_needs_review else "not_needs_review"
    selected = [q for q in questions if gate_ok(q, gate)]
    if not selected:
        raise ValueError(
            "출력할 문항이 없습니다. 게이트를 확인하세요 "
            "(--gate approved|not_needs_review|all)."
        )
    selected.sort(key=lambda q: q.get("question_number", 0))
    objective = [q for q in selected if not is_essay(q)]
    essay = [q for q in selected if is_essay(q)]

    doc = _render_document(exam, objective, essay, include_explanation, group_by)

    out = Path(out_path) if out_path else Path(exam_json_path).with_suffix(".cbt.docx")
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))

    return {
        "export_type": "cbt_docx_template",
        "output": str(out),
        "objective_count": len(objective),
        "essay_count": len(essay),
        "gate": gate,
        "group_by": group_by,
        "hwp_workflow": "한글에서 .docx 열기 → 필요 시 편집 → 다른 이름으로 저장(.hwp)",
    }


def main() -> int:
    ap = argparse.ArgumentParser(description="의학교육실 과정시험 HWP 양식 .docx 익스포트")
    ap.add_argument("exam_json")
    ap.add_argument("--out", default=None)
    ap.add_argument("--include-needs-review", action="store_true",
                    help="검토중 문항까지 포함(=--gate all). 미리보기용")
    ap.add_argument("--gate", default=None,
                    choices=["approved", "not_needs_review", "all"],
                    help="추출 게이트. 기본 not_needs_review, 운영 제출은 approved")
    ap.add_argument("--no-explanation", action="store_true")
    ap.add_argument("--group-by", default="course_name", choices=["course_name", "none"])
    a = ap.parse_args()
    result = build_cbt_docx_export(
        a.exam_json,
        a.out,
        include_needs_review=a.include_needs_review,
        include_explanation=not a.no_explanation,
        group_by=a.group_by,
        gate=a.gate,
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
