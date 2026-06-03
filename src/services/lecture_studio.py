from __future__ import annotations

import json
import os
import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from urllib.parse import quote

import requests

from scripts.generate_lecture_questions import (
    PIPELINE_VERSION,
    build_generation_prompt,
    extract_json_payload,
    extract_text,
    generate_gemini,
    generate_openai,
    normalize_questions,
    slugify,
)


DATA_ROOT = Path("data_private/studio")
UPLOAD_DIR = DATA_ROOT / "uploads"
EXTRACTED_DIR = DATA_ROOT / "extracted"
GENERATED_DIR = DATA_ROOT / "generated"
IMAGE_DIR = DATA_ROOT / "images"
CONVERTED_DIR = DATA_ROOT / "converted"
MEDIA_DIR = DATA_ROOT / "media_bank"
MEDIA_ASSET_DIR = MEDIA_DIR / "assets"
MEDIA_INDEX_PATH = MEDIA_DIR / "media_assets.json"
QUESTION_BANK_DIR = DATA_ROOT / "question_bank"
REVIEW_SET_DIR = DATA_ROOT / "review_sets"
EXPORT_SET_DIR = DATA_ROOT / "export_sets"

MODEL_CATALOG: dict[str, list[dict[str, str]]] = {
    "claude-cli": [
        {"id": "sonnet", "label": "Claude Code Sonnet latest", "note": "현재 로그인된 Claude Code 계정 사용"},
        {"id": "opus", "label": "Claude Code Opus latest", "note": "현재 로그인된 Claude Code 계정 사용"},
        {"id": "claude-sonnet-4-6", "label": "Claude Sonnet 4.6", "note": "Claude Code full model ID"},
        {"id": "claude-opus-4-8", "label": "Claude Opus 4.8", "note": "Claude Code full model ID"},
    ],
    "anthropic": [
        {"id": "claude-sonnet-4-6", "label": "Claude Sonnet 4.6", "note": "ANTHROPIC_API_KEY 필요"},
        {"id": "claude-opus-4-8", "label": "Claude Opus 4.8", "note": "ANTHROPIC_API_KEY 필요"},
        {"id": "claude-haiku-4-5-20251001", "label": "Claude Haiku 4.5", "note": "ANTHROPIC_API_KEY 필요"},
    ],
    "openai": [
        {"id": "gpt-5.2", "label": "GPT-5.2", "note": "OPENAI_API_KEY 필요"},
        {"id": "gpt-5.2-chat-latest", "label": "GPT-5.2 Chat latest", "note": "OPENAI_API_KEY 필요"},
        {"id": "gpt-5.1", "label": "GPT-5.1", "note": "OPENAI_API_KEY 필요"},
        {"id": "gpt-5", "label": "GPT-5", "note": "OPENAI_API_KEY 필요"},
        {"id": "gpt-5-mini", "label": "GPT-5 mini", "note": "OPENAI_API_KEY 필요"},
        {"id": "gpt-5-nano", "label": "GPT-5 nano", "note": "OPENAI_API_KEY 필요"},
        {"id": "gpt-4.1", "label": "GPT-4.1", "note": "OPENAI_API_KEY 필요"},
        {"id": "gpt-4o-mini", "label": "GPT-4o mini", "note": "저비용 fallback"},
    ],
    "gemini": [
        {"id": "gemini-2.5-flash", "label": "Gemini 2.5 Flash", "note": "GEMINI_API_KEY 필요"},
        {"id": "gemini-2.5-pro", "label": "Gemini 2.5 Pro", "note": "GEMINI_API_KEY 필요"},
    ],
    "prompt-only": [
        {"id": "prompt-only", "label": "프롬프트만 생성", "note": "API 키 없이 추출/프롬프트 저장"},
    ],
}


@dataclass
class SavedUpload:
    path: Path
    original_name: str
    kind: str


def ensure_studio_dirs() -> None:
    for directory in (
        UPLOAD_DIR,
        EXTRACTED_DIR,
        GENERATED_DIR,
        IMAGE_DIR,
        CONVERTED_DIR,
        MEDIA_DIR,
        MEDIA_ASSET_DIR,
        QUESTION_BANK_DIR,
        REVIEW_SET_DIR,
        EXPORT_SET_DIR,
    ):
        directory.mkdir(parents=True, exist_ok=True)


def read_json_list(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return []
    return data if isinstance(data, list) else []


def write_text_atomic(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        "w",
        encoding="utf-8",
        dir=str(path.parent),
        delete=False,
    ) as temp_file:
        temp_file.write(content)
        temp_name = temp_file.name
    os.replace(temp_name, path)


def write_json_list(path: Path, rows: list[dict[str, Any]]) -> None:
    write_text_atomic(path, json.dumps(rows, ensure_ascii=False, indent=2))


def split_metadata_values(value: str | None) -> list[str]:
    if not value:
        return []
    return [
        item.strip()
        for item in re.split(r"[,;\n]+", str(value))
        if item.strip()
    ]


def claude_cli_available() -> bool:
    return bool(shutil.which("claude"))


def get_model_catalog() -> dict[str, Any]:
    return {
        "providers": [
            {
                "id": "auto",
                "label": "자동 선택",
                "available": True,
                "models": [{"id": "auto", "label": "자동", "note": "사용 가능한 계정/키를 순서대로 사용"}],
            },
            {
                "id": "claude-cli",
                "label": "Claude Code 계정",
                "available": claude_cli_available(),
                "models": MODEL_CATALOG["claude-cli"],
            },
            {
                "id": "anthropic",
                "label": "Claude API",
                "available": bool(os.getenv("ANTHROPIC_API_KEY")),
                "models": MODEL_CATALOG["anthropic"],
            },
            {
                "id": "openai",
                "label": "OpenAI API",
                "available": bool(os.getenv("OPENAI_API_KEY")),
                "models": MODEL_CATALOG["openai"],
            },
            {
                "id": "gemini",
                "label": "Google Gemini",
                "available": bool(os.getenv("GEMINI_API_KEY")),
                "models": MODEL_CATALOG["gemini"],
            },
            {
                "id": "prompt-only",
                "label": "프롬프트만 생성",
                "available": True,
                "models": MODEL_CATALOG["prompt-only"],
            },
        ]
    }


def resolve_studio_provider(provider: str) -> str:
    if provider and provider != "auto":
        return provider
    if os.getenv("OPENAI_API_KEY"):
        return "openai"
    if os.getenv("ANTHROPIC_API_KEY"):
        return "anthropic"
    if os.getenv("GEMINI_API_KEY"):
        return "gemini"
    if claude_cli_available():
        return "claude-cli"
    return "prompt-only"


def default_model_for_provider(provider: str, selected_model: str | None = None) -> str:
    if selected_model and selected_model != "auto":
        return selected_model
    defaults = {
        "claude-cli": "sonnet",
        "anthropic": "claude-sonnet-4-6",
        "openai": os.getenv("OPENAI_MODEL", "gpt-5.2"),
        "gemini": os.getenv("GEMINI_MODEL", "gemini-2.5-flash"),
        "prompt-only": "prompt-only",
    }
    return defaults.get(provider, "prompt-only")


def timestamp_slug() -> str:
    return datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")


def save_upload_bytes(content: bytes, filename: str, *, kind: str) -> SavedUpload:
    ensure_studio_dirs()
    safe_stem = slugify(Path(filename or f"{kind}.txt").stem)
    suffix = Path(filename or "").suffix.lower() or ".txt"
    path = UPLOAD_DIR / f"{timestamp_slug()}_{kind}_{safe_stem}{suffix}"
    path.write_bytes(content or b"")
    return SavedUpload(path=path, original_name=filename or path.name, kind=kind)


def save_media_asset(
    content: bytes,
    filename: str,
    *,
    asset_type: str = "clinical_photo",
    modality: str = "",
    subject: str = "",
    unit: str = "",
    diagnosis: str = "",
    caption: str = "",
    key_findings: str = "",
    deidentified: bool = False,
    approved_for_question_use: bool = False,
    faculty_note: str = "",
) -> dict[str, Any]:
    ensure_studio_dirs()
    safe_stem = slugify(Path(filename or "media").stem)
    suffix = Path(filename or "").suffix.lower() or ".png"
    asset_id = f"media_{timestamp_slug()}_{safe_stem}"
    stored_name = f"{asset_id}{suffix}"
    asset_path = MEDIA_ASSET_DIR / stored_name
    asset_path.write_bytes(content or b"")
    now = datetime.now(timezone.utc).isoformat()
    asset = {
        "asset_id": asset_id,
        "asset_type": asset_type or "clinical_photo",
        "modality": modality,
        "subject": subject,
        "unit": unit,
        "diagnosis": diagnosis,
        "caption": caption,
        "key_findings": split_metadata_values(key_findings),
        "deidentified": bool(deidentified),
        "approved_for_question_use": bool(approved_for_question_use),
        "faculty_note": faculty_note,
        "original_name": filename or stored_name,
        "stored_name": stored_name,
        "file_path": str(asset_path),
        "url": f"/api/media/assets/{quote(stored_name)}",
        "review_status": "approved" if approved_for_question_use else "needs_review",
        "created_at": now,
        "updated_at": now,
    }
    assets = read_json_list(MEDIA_INDEX_PATH)
    assets.append(asset)
    write_json_list(MEDIA_INDEX_PATH, assets)
    return asset


def list_media_assets() -> list[dict[str, Any]]:
    ensure_studio_dirs()
    assets = read_json_list(MEDIA_INDEX_PATH)
    return sorted(assets, key=lambda item: str(item.get("created_at", "")), reverse=True)


def delete_media_asset(asset_id: str) -> dict[str, Any]:
    ensure_studio_dirs()
    safe_id = str(asset_id or "").strip()
    assets = read_json_list(MEDIA_INDEX_PATH)
    kept: list[dict[str, Any]] = []
    deleted: dict[str, Any] | None = None
    for asset in assets:
        if str(asset.get("asset_id")) == safe_id:
            deleted = asset
            continue
        kept.append(asset)
    if deleted is None:
        raise FileNotFoundError(safe_id)

    stored_name = Path(str(deleted.get("stored_name") or "")).name
    if stored_name:
        asset_path = (MEDIA_ASSET_DIR / stored_name).resolve()
        asset_root = MEDIA_ASSET_DIR.resolve()
        if asset_root in asset_path.parents or asset_path == asset_root:
            asset_path.unlink(missing_ok=True)
    write_json_list(MEDIA_INDEX_PATH, kept)
    return deleted


def get_media_assets(asset_ids: list[str]) -> list[dict[str, Any]]:
    wanted = {asset_id for asset_id in asset_ids if asset_id}
    return [asset for asset in list_media_assets() if asset.get("asset_id") in wanted]


def extract_saved_upload(upload: SavedUpload) -> str:
    return extract_text(upload.path)


def build_context_block(title: str, uploads: list[SavedUpload], *, max_chars_per_file: int, max_total_chars: int) -> str:
    if not uploads:
        return ""

    blocks: list[str] = []
    used = 0
    for upload in uploads:
        try:
            text = extract_saved_upload(upload)
        except Exception as exc:
            text = f"[추출 실패: {exc}]"

        block = f"--- {title}: {upload.original_name} ---\n{text[:max_chars_per_file]}"
        remaining = max_total_chars - used
        if remaining <= 0:
            break
        if len(block) > remaining:
            block = block[:remaining].rstrip() + "\n[TRUNCATED]"
        blocks.append(block)
        used += len(block)

    return "\n\n".join(blocks).strip()


def append_style_context(prompt: str, style_context: str) -> str:
    if not style_context:
        return prompt
    return (
        f"{prompt}\n\n"
        "[기출문항 유형 참고]\n"
        "아래 자료는 문항의 문장 길이, 선지 구성, 난이도, 함정 유형을 참고하기 위한 자료입니다. "
        "새 문항의 의학 내용은 반드시 강의록과 승인 근거자료 범위에서만 생성하세요.\n"
        f"{style_context}"
    )


def append_question_type_instruction(prompt: str, question_type: str, visual_candidates: list[dict[str, Any]]) -> str:
    type_map = {
        "clinical_case": "임상 증례형: 환자 정보, 증상, 검사 소견을 바탕으로 진단/치료/다음 단계를 묻습니다.",
        "image_based": "이미지/자료해석형: 제공된 강의자료/근거자료의 그림, 표, 영상, 검사자료를 해석하는 문항을 우선 생성합니다.",
        "basic_concept": "개념확인형: 강의 핵심 개념, 정의, 분류, 기준을 묻되 단순 암기보다 적용형으로 만듭니다.",
        "mechanism": "병태생리/기전형: 질환 발생 기전, 약물 작용기전, 생리학적 원리를 묻습니다.",
        "diagnostic": "진단/검사 선택형: 감별진단, 초기검사, 확진검사, 검사 결과 해석을 묻습니다.",
        "management": "치료/처치 선택형: 치료 원칙, 금기, 다음 처치, 추적관찰을 묻습니다.",
        "mixed": "혼합형: 증례형, 개념형, 진단/치료형을 균형 있게 섞습니다.",
    }
    instruction = type_map.get(question_type, type_map["clinical_case"])
    visual_lines = "\n".join(
        (
            f"- {candidate.get('id')}: {candidate.get('source_name')} "
            f"{candidate.get('locator_label') or 'p.' + str(candidate.get('page', '-'))}. "
            f"텍스트 단서: {str(candidate.get('text_preview') or '')[:140]}"
        )
        for candidate in visual_candidates[:12]
    ) or "- 없음"
    return f"""
{prompt}

[문제 유형 지침]
- 선택된 문제 유형: {question_type}
- 작성 방향: {instruction}
- question_type 필드에는 "{question_type}"를 넣으세요.
- 이미지/자료해석형을 선택한 경우, 문제 본문에는 "[첨부 이미지 참조]" 같은 플레이스홀더를 쓰지 마세요. 실제 이미지는 시스템이 image_refs로 문항 본문 안에 삽입합니다.
- 이미지/자료해석형을 선택한 경우, pma_solution.source_anchor에는 반드시 관련 자료명, slide/page 번호, 이미지 후보 ID 중 하나 이상을 남기세요.
- 이미지 후보 중 환자 사진, 신체진찰 사진, X-ray/CXR, CT, MRI, 초음파, 심전도/검사 영상이 있으면 이를 최우선으로 활용하세요.
- 표, 프로토콜, 요약 슬라이드만 있는 경우에는 자료해석 가치가 명확할 때만 이미지형 문항으로 활용하고, 없는 환자 사진이나 영상검사를 지어내지 마세요.

[사용 가능한 이미지 후보]
{visual_lines}
""".strip()


def append_reference_instruction(
    prompt: str,
    lecture: SavedUpload,
    evidence_uploads: list[SavedUpload],
    *,
    reference_policy: str,
) -> str:
    evidence_names = [upload.original_name for upload in evidence_uploads]
    evidence_lines = "\n".join(f"- {name}" for name in evidence_names) or "- 없음"
    if reference_policy == "local_open":
        policy_instruction = """
- 로컬 개인 학습/검수용 초안이므로, 업로드 자료 외에도 모델이 알고 있는 표준 의학지식, 교과서, 진료지침, 국가고시/KMLE/USMLE식 지식을 보조적으로 참고해도 됩니다.
- Harrison, 홍창의 소아과학, Nelson Textbook of Pediatrics, UpToDate, 학회 guideline, case report 등은 실제 원문을 인용하지 말고 참고 범주/근거 요약 수준으로만 적으세요.
- 업로드되지 않은 reference를 source에 적는 경우 basis 끝에 "(모델 일반지식 기반, 원문 검수 필요)"를 붙이세요.
- 외부 reference는 정확한 판/쪽수/문구를 지어내지 마세요.
""".strip()
    else:
        policy_instruction = """
- source에는 실제 제공된 자료명만 적습니다.
- Harrison, 홍창의 소아과학, Guideline, case report 등은 사용자가 승인 근거자료로 업로드한 경우에만 source에 적습니다.
- 업로드되지 않은 교과서/가이드라인/논문을 참고했다고 쓰지 마세요.
- 외부 근거가 필요하지만 자료가 없으면 needs_review=true 및 review_reasons에 "reference_needed"를 포함하세요.
""".strip()
    return f"""
{prompt}

[참고문헌/근거 기록 지침]
- 각 문항에는 reference_notes 배열을 반드시 포함하세요.
- reference_notes는 ref_no, source, basis 필드를 가진 객체 배열입니다.
- ref_no는 1부터 시작하고 문항별로 한 줄에 하나씩 번호를 붙입니다.
- 강의자료 근거는 "{lecture.original_name}"로 적습니다.
- explanation 마지막에는 "참고문헌:" 아래에 1., 2., 3. 형식으로 각 줄 번호를 나누어 적으세요.

{policy_instruction}

[이번 요청에서 제공된 승인 근거자료]
{evidence_lines}
""".strip()


def append_media_instruction(
    prompt: str,
    visual_candidates: list[dict[str, Any]],
    *,
    selected_media_assets: list[dict[str, Any]],
    image_description: str = "",
) -> str:
    if not visual_candidates and not image_description.strip():
        return prompt

    selected_lines = []
    for asset in selected_media_assets:
        findings = asset.get("key_findings") if isinstance(asset.get("key_findings"), list) else []
        selected_lines.append(
            "- "
            f"{asset.get('asset_id')}: {asset.get('modality') or asset.get('asset_type')} | "
            f"진단/주제: {asset.get('diagnosis') or '-'} | "
            f"핵심 소견: {', '.join(str(item) for item in findings) or '-'} | "
            f"설명: {asset.get('caption') or '-'}"
        )
    selected_block = "\n".join(selected_lines) or "- 선택된 Media Bank 이미지 없음"
    description_block = image_description.strip() or "- 별도 이미지 설명 없음"

    return f"""
{prompt}

[이미지 기반 문항 제작 지침]
- 교수자가 선택한 Media Bank 자료 또는 아래 이미지 설명을 문항의 제시자료로 사용하세요.
- 문제 본문은 텍스트 임상 상황을 먼저 제시하고, 이미지 소견을 함께 종합해 풀 수 있게 작성하세요.
- 선지는 이미지 없이도 대충 맞힐 수 있는 단순 암기형이 아니라, 이미지 소견과 임상 정보가 함께 필요하도록 만드세요.
- 해설에는 image_findings 또는 pma_solution.reasoning_summary에 이미지에서 확인해야 할 핵심 소견을 명시하세요.
- 원자료가 실제 환자/영상/병리 자료라면 비식별화와 교수 검수가 필요하다는 점을 review_reasons에 반영하세요.

[교수 선택 Media Bank 자료]
{selected_block}

[교수 입력 이미지 설명]
{description_block}
""".strip()


def append_table_instruction(prompt: str) -> str:
    return f"""
{prompt}

[표/검사자료 생성 지침]
- 필요하면 문항에 사용할 짧은 표를 직접 생성해도 됩니다.
- 표는 활력징후, CBC/chemistry/ABGA, CSF, urinalysis, ECG 판독 요약, 치료 전후 변화처럼 문항 해결에 실제로 필요한 자료일 때만 사용하세요.
- 표가 들어가는 문항은 problem 본문에 "다음 표를 참고하라"처럼 자연스럽게 연결하세요.
- JSON에는 선택적으로 data_table 필드를 포함할 수 있습니다.
- data_table 형식:
  {{
    "title": "표 제목",
    "columns": ["항목", "결과", "참고치"],
    "rows": [
      ["WBC", "18,500/µL", "4,000-10,000/µL"],
      ["CRP", "12 mg/dL", "<0.5 mg/dL"]
    ]
  }}
- 표 안의 수치와 단위는 의학적으로 그럴듯해야 하며, 강의자료 근거가 부족하면 needs_review=true와 review_reasons에 "table_value_review_needed"를 포함하세요.
""".strip()


def write_json(path: Path, payload: Any) -> None:
    write_text_atomic(path, json.dumps(payload, ensure_ascii=False, indent=2))


def archive_question_set(source_slug: str, metadata: dict[str, Any], records: list[dict[str, Any]]) -> Path:
    ensure_studio_dirs()
    archive_path = QUESTION_BANK_DIR / f"{source_slug}.question_set.json"
    now = datetime.now(timezone.utc).isoformat()
    packet = {
        "set_id": source_slug,
        "review_status": "faculty_review_pending",
        "created_at": now,
        "updated_at": now,
        "metadata": metadata,
        "questions": records,
        "review_events": [],
    }
    write_json(archive_path, packet)
    review_path = REVIEW_SET_DIR / f"{source_slug}.review_set.json"
    write_json(
        review_path,
        {
            "set_id": source_slug,
            "review_status": "faculty_review_pending",
            "created_at": packet["created_at"],
            "updated_at": packet["updated_at"],
            "source_question_set": str(archive_path),
            "question_count": len(records),
            "needs_review_count": sum(1 for record in records if record.get("needs_review")),
        },
    )
    return archive_path


def safe_question_set_path(set_id: str) -> Path:
    safe_id = Path(str(set_id or "")).name
    if not safe_id or safe_id != str(set_id):
        raise FileNotFoundError("invalid set_id")
    path = QUESTION_BANK_DIR / f"{safe_id}.question_set.json"
    if not path.exists():
        raise FileNotFoundError(set_id)
    return path


def summarize_question_set(packet: dict[str, Any], *, path: Path | None = None) -> dict[str, Any]:
    questions = packet.get("questions") if isinstance(packet.get("questions"), list) else []
    metadata = packet.get("metadata") if isinstance(packet.get("metadata"), dict) else {}
    return {
        "set_id": packet.get("set_id") or (path.stem.replace(".question_set", "") if path else ""),
        "review_status": packet.get("review_status", "draft"),
        "created_at": packet.get("created_at"),
        "updated_at": packet.get("updated_at"),
        "source_name": metadata.get("source_name"),
        "subject": metadata.get("subject"),
        "unit": metadata.get("unit"),
        "provider": metadata.get("provider"),
        "model": metadata.get("model"),
        "question_type": metadata.get("question_type"),
        "question_count": len(questions),
        "image_question_count": sum(1 for question in questions if question.get("image_refs")),
        "needs_review_count": sum(1 for question in questions if question.get("needs_review")),
        "approved_count": sum(1 for question in questions if question.get("review_status") == "approved"),
        "rejected_count": sum(1 for question in questions if question.get("review_status") == "rejected"),
        "path": str(path) if path else None,
    }


def list_question_sets(limit: int = 20) -> list[dict[str, Any]]:
    ensure_studio_dirs()
    rows: list[dict[str, Any]] = []
    for path in sorted(QUESTION_BANK_DIR.glob("*.question_set.json"), key=lambda item: item.stat().st_mtime, reverse=True):
        try:
            packet = json.loads(path.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            continue
        rows.append(summarize_question_set(packet, path=path))
        if len(rows) >= limit:
            break
    return rows


def load_question_set(set_id: str, *, include_summary: bool = True) -> dict[str, Any]:
    ensure_studio_dirs()
    path = safe_question_set_path(set_id)
    try:
        packet = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise ValueError(f"문항 세트 JSON을 읽을 수 없습니다: {set_id}") from exc
    if not isinstance(packet, dict):
        raise ValueError(f"문항 세트 형식이 올바르지 않습니다: {set_id}")
    packet.setdefault("set_id", Path(path.name).name.replace(".question_set.json", ""))
    packet.setdefault("questions", [])
    if include_summary:
        packet["summary"] = summarize_question_set(packet, path=path)
    return packet


def normalize_review_reasons(value: Any) -> list[str]:
    if isinstance(value, list):
        items = value
    elif isinstance(value, str):
        items = re.split(r"[,;\n]+", value)
    else:
        items = []
    return sorted({str(item).strip() for item in items if str(item).strip()})


def normalize_review_status(value: str | None) -> str:
    status = str(value or "draft").strip()
    allowed = {"draft", "needs_revision", "approved", "rejected", "exported"}
    return status if status in allowed else "draft"


def recompute_set_review_status(questions: list[dict[str, Any]]) -> str:
    if not questions:
        return "draft"
    statuses = [normalize_review_status(question.get("review_status")) for question in questions]
    if statuses and all(status == "draft" for status in statuses):
        return "draft"
    if statuses and all(status == "approved" for status in statuses):
        return "approved"
    if any(status == "rejected" for status in statuses):
        return "faculty_review_pending"
    return "faculty_review_pending"


def sync_review_set_summary(packet: dict[str, Any], archive_path: Path) -> None:
    if not packet.get("set_id"):
        raise ValueError("문항 세트에 set_id가 없습니다.")
    review_path = REVIEW_SET_DIR / f"{packet.get('set_id')}.review_set.json"
    summary = summarize_question_set(packet, path=archive_path)
    write_json(
        review_path,
        {
            "set_id": summary["set_id"],
            "review_status": packet.get("review_status", "faculty_review_pending"),
            "created_at": packet.get("created_at"),
            "updated_at": packet.get("updated_at"),
            "source_question_set": str(archive_path),
            "question_count": summary["question_count"],
            "needs_review_count": summary["needs_review_count"],
            "approved_count": summary["approved_count"],
            "rejected_count": summary["rejected_count"],
        },
    )


def save_question_set_packet(packet: dict[str, Any], path: Path) -> dict[str, Any]:
    now = datetime.now(timezone.utc).isoformat()
    packet.pop("summary", None)
    questions = packet.get("questions") if isinstance(packet.get("questions"), list) else []
    packet["questions"] = questions
    packet["review_status"] = recompute_set_review_status(questions)
    packet["updated_at"] = now
    write_json(path, packet)
    sync_review_set_summary(packet, path)
    response_packet = dict(packet)
    response_packet["summary"] = summarize_question_set(packet, path=path)
    return response_packet


def find_question_index(packet: dict[str, Any], question_id: str) -> int:
    questions = packet.get("questions") if isinstance(packet.get("questions"), list) else []
    for index, question in enumerate(questions):
        if str(question.get("question_id")) == str(question_id):
            return index
    raise KeyError(question_id)


def add_review_event(
    packet: dict[str, Any],
    *,
    question_id: str,
    event_type: str,
    before: dict[str, Any] | None = None,
    after: dict[str, Any] | None = None,
    comment: str = "",
    actor_id: str = "local_faculty",
) -> None:
    events = packet.get("review_events")
    if not isinstance(events, list):
        events = []
        packet["review_events"] = events
    now = datetime.now(timezone.utc).isoformat()
    events.append(
        {
            "event_id": f"review_{timestamp_slug()}_{len(events) + 1:04d}",
            "set_id": packet.get("set_id"),
            "question_id": question_id,
            "actor_id": actor_id,
            "event_type": event_type,
            "before": before or {},
            "after": after or {},
            "comment": comment,
            "created_at": now,
        }
    )


def update_question_review(
    set_id: str,
    question_id: str,
    updates: dict[str, Any],
    *,
    action: str = "edited",
    actor_id: str = "local_faculty",
    comment: str = "",
) -> dict[str, Any]:
    path = safe_question_set_path(set_id)
    packet = load_question_set(set_id, include_summary=False)
    questions = packet.get("questions") if isinstance(packet.get("questions"), list) else []
    index = find_question_index(packet, question_id)
    before = dict(questions[index])
    question = dict(questions[index])

    editable_fields = {
        "problem",
        "options",
        "answer",
        "explanation",
        "review_status",
        "needs_review",
        "review_reasons",
        "reference_notes",
        "image_refs",
        "image_match_confidence",
    }
    for key, value in updates.items():
        if key not in editable_fields:
            continue
        if key == "options":
            options = [str(item).strip() for item in list(value or []) if str(item).strip()]
            while len(options) < 5:
                options.append(f"검수 필요 보기 {len(options) + 1}")
            question[key] = options[:5]
        elif key == "answer":
            try:
                answer = int(value)
            except (TypeError, ValueError):
                answer = question.get("answer") or 1
            question[key] = max(1, min(5, answer))
        elif key == "needs_review":
            question[key] = bool(value)
        elif key == "review_reasons":
            question[key] = normalize_review_reasons(value)
        elif key == "review_status":
            question[key] = normalize_review_status(str(value))
        else:
            question[key] = value

    now = datetime.now(timezone.utc).isoformat()
    question["updated_at"] = now
    if action == "approved":
        question["review_status"] = "approved"
        question["needs_review"] = False
        question["review_reasons"] = []
        question["approved_at"] = now
    elif action == "rejected":
        question["review_status"] = "rejected"
        question["needs_review"] = True
        reasons = normalize_review_reasons(question.get("review_reasons"))
        if "faculty_rejected" not in reasons:
            reasons.append("faculty_rejected")
        question["review_reasons"] = sorted(set(reasons))
        question["rejected_at"] = now
    elif question.get("review_status") == "approved":
        question["needs_review"] = False
        question["review_reasons"] = []
    elif question.get("review_status") in {"needs_revision", "rejected"}:
        question["needs_review"] = True

    questions[index] = question
    packet["questions"] = questions
    add_review_event(
        packet,
        question_id=question_id,
        event_type=action,
        before=before,
        after=question,
        comment=comment,
        actor_id=actor_id,
    )
    packet = save_question_set_packet(packet, path)
    return {
        "set": packet,
        "summary": packet["summary"],
        "question": question,
    }


def parse_numbered_sections(text: str, marker: str) -> dict[int, str]:
    matches = list(re.finditer(rf"(?m)^=== {re.escape(marker)} (\d+) ===\s*$", str(text or "")))
    if not matches:
        return {1: str(text or "")}

    pages: dict[int, str] = {}
    for idx, match in enumerate(matches):
        page_num = int(match.group(1))
        start = match.end()
        end = matches[idx + 1].start() if idx + 1 < len(matches) else len(text)
        pages[page_num] = text[start:end].strip()
    return pages


def parse_page_sections(text: str) -> dict[int, str]:
    return parse_numbered_sections(text, "page")


def tokenize_for_match(text: str) -> set[str]:
    tokens = re.findall(r"[A-Za-z][A-Za-z0-9-]{2,}|[가-힣]{2,}", str(text or "").lower())
    stopwords = {
        "문항", "문제", "정답", "해설", "가장", "다음", "대한", "있는", "경우", "환자",
        "강의록", "근거", "출제", "포인트", "검사", "진단", "치료", "소견",
        "patient", "most", "likely", "following", "diagnosis", "management",
    }
    return {token for token in tokens if token not in stopwords and len(token) >= 2}


def visual_candidate_priority(candidate: dict[str, Any]) -> float:
    text = " ".join(
        [
            str(candidate.get("id") or ""),
            str(candidate.get("source_name") or ""),
            str(candidate.get("locator_label") or ""),
            str(candidate.get("text_preview") or ""),
        ]
    ).lower()
    positive_groups = [
        (
            0.55,
            [
                "환자 사진",
                "patient photo",
                "clinical photo",
                "gross",
                "skin lesion",
                "rash",
                "소견 사진",
                "진찰 사진",
            ],
        ),
        (
            0.5,
            [
                "ct",
                "mri",
                "초음파",
                "ultrasound",
                "sonography",
                "x-ray",
                "xray",
                "x ray",
                "cxr",
                "radiograph",
                "방사선",
                "영상",
                "흉부 사진",
                "단순촬영",
            ],
        ),
        (
            0.38,
            [
                "ecg",
                "ekg",
                "심전도",
                "pbs",
                "peripheral blood smear",
                "blood smear",
                "말초혈액",
                "도말",
                "검사 사진",
                "검사 영상",
                "pathology",
                "병리",
                "조직",
                "조직병리",
                "병리 슬라이드",
                "현미경",
                "microscopy",
                "histology",
                "biopsy",
            ],
        ),
        (
            0.32,
            [
                "ballard",
                "각창",
                "square window",
                "오금",
                "popliteal",
                "스카프",
                "scarf",
                "heel to ear",
                "발뒤꿈치",
                "kangaroo",
                "head hood",
                "cpap",
                "c-pap",
            ],
        ),
        (
            0.2,
            [
                "사진",
                "image",
                "figure",
                "그림",
                "도식",
            ],
        ),
    ]
    negative_terms = [
        "목차",
        "학습목표",
        "요약",
        "reference",
        "참고문헌",
        "프로토콜",
        "protocol",
        "algorithm",
        "표 ",
        "table",
        "통계",
        "빈도",
        "그래프",
        "graph",
        "chart",
    ]

    def has_keyword(keyword: str) -> bool:
        key = keyword.lower()
        if re.fullmatch(r"[a-z0-9]{1,4}", key):
            return bool(re.search(rf"(?<![a-z0-9]){re.escape(key)}(?![a-z0-9])", text))
        return key in text

    picture_count = int(candidate.get("picture_count") or 0)
    picture_ratio = float(candidate.get("largest_picture_area_ratio") or 0)
    priority = 0.25 if candidate.get("type") == "uploaded_image" else 0.02
    if picture_count:
        priority = max(priority, 0.22)
    if picture_ratio >= 0.18:
        priority = max(priority, 0.28)
    for weight, keywords in positive_groups:
        if any(has_keyword(keyword) for keyword in keywords):
            priority = max(priority, weight)
    if any(has_keyword(term) for term in negative_terms):
        priority = max(0.0, priority - 0.12)
    return round(priority, 3)


def extract_pdf_snapshot_candidates(
    pdf_path: Path,
    upload: SavedUpload,
    source_slug: str,
    source_text: str,
    *,
    max_pages: int = 24,
    marker: str = "page",
    candidate_type: str = "page_snapshot",
    media_hints: dict[int, dict[str, Any]] | None = None,
) -> list[dict[str, Any]]:
    try:
        import fitz
    except Exception:
        return []

    page_texts = parse_numbered_sections(source_text, marker)
    candidates: list[dict[str, Any]] = []
    doc = fitz.open(str(pdf_path))
    locator_prefix = "slide " if marker == "slide" else "p."
    try:
        for page_index, page in enumerate(doc):
            if page_index >= max_pages:
                break
            page_num = page_index + 1
            media_hint = (media_hints or {}).get(page_num, {})
            filename = f"{source_slug}_p{page_num:03d}.png"
            image_path = IMAGE_DIR / filename
            if not image_path.exists():
                pix = page.get_pixmap(matrix=fitz.Matrix(1.6, 1.6), alpha=False)
                pix.save(str(image_path))

            page_text = page_texts.get(page_num, "")
            candidate = {
                "id": f"{source_slug}_p{page_num:03d}",
                "type": candidate_type,
                "source_name": upload.original_name,
                "page": page_num,
                "locator_label": f"{locator_prefix}{page_num}",
                "filename": filename,
                "url": f"/api/studio/images/{quote(filename)}",
                "text_preview": page_text[:220],
                "picture_count": media_hint.get("picture_count", 0),
                "largest_picture_area_ratio": media_hint.get("largest_picture_area_ratio", 0),
                "_tokens": tokenize_for_match(page_text),
            }
            candidate["visual_priority"] = visual_candidate_priority(candidate)
            candidates.append(candidate)
    finally:
        doc.close()
    return candidates


def office_converter_path() -> str | None:
    for command in ("soffice", "libreoffice", "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        found = shutil.which(command) if not command.startswith("/") else command
        if found and Path(found).exists():
            return found
    return None


def convert_office_to_pdf(path: Path, source_slug: str) -> Path | None:
    converter = office_converter_path()
    if not converter:
        return None

    outdir = CONVERTED_DIR / source_slug
    outdir.mkdir(parents=True, exist_ok=True)
    existing = list(outdir.glob("*.pdf"))
    if existing:
        return max(existing, key=lambda item: item.stat().st_mtime)

    with tempfile.TemporaryDirectory(dir=str(CONVERTED_DIR)) as profile_dir:
        profile_uri = Path(profile_dir).resolve().as_uri()
        cmd = [
            converter,
            "--headless",
            f"-env:UserInstallation={profile_uri}",
            "--convert-to",
            "pdf",
            "--outdir",
            str(outdir),
            str(path),
        ]
        result = subprocess.run(cmd, capture_output=True, text=True, timeout=120)
    if result.returncode != 0:
        return None
    converted = list(outdir.glob("*.pdf"))
    return max(converted, key=lambda item: item.stat().st_mtime) if converted else None


def extract_pptx_media_hints(path: Path) -> dict[int, dict[str, Any]]:
    try:
        from pptx import Presentation
    except Exception:
        return {}

    hints: dict[int, dict[str, Any]] = {}
    prs = Presentation(str(path))
    slide_area = max(int(prs.slide_width) * int(prs.slide_height), 1)
    for slide_index, slide in enumerate(prs.slides, 1):
        picture_areas: list[int] = []
        for shape in slide.shapes:
            if getattr(shape, "shape_type", None) == 13:
                picture_areas.append(int(shape.width) * int(shape.height))
        if picture_areas:
            hints[slide_index] = {
                "picture_count": len(picture_areas),
                "largest_picture_area_ratio": round(max(picture_areas) / slide_area, 3),
            }
    return hints


def extract_visual_candidates(upload: SavedUpload, source_slug: str, source_text: str, *, max_pages: int = 24) -> list[dict[str, Any]]:
    suffix = upload.path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf_snapshot_candidates(
            upload.path,
            upload,
            source_slug,
            source_text,
            max_pages=max_pages,
            marker="page",
            candidate_type="page_snapshot",
        )
    if suffix == ".pptx":
        converted_pdf = convert_office_to_pdf(upload.path, source_slug)
        if not converted_pdf:
            return []
        media_hints = extract_pptx_media_hints(upload.path)
        return extract_pdf_snapshot_candidates(
            converted_pdf,
            upload,
            source_slug,
            source_text,
            max_pages=max_pages,
            marker="slide",
            candidate_type="slide_snapshot",
            media_hints=media_hints,
        )
    return []


def direct_image_candidate(upload: SavedUpload, source_slug: str) -> dict[str, Any] | None:
    if upload.path.suffix.lower() not in {".png", ".jpg", ".jpeg", ".webp"}:
        return None
    filename = f"{source_slug}_{slugify(Path(upload.original_name).stem)}{upload.path.suffix.lower()}"
    image_path = IMAGE_DIR / filename
    if not image_path.exists():
        shutil.copyfile(upload.path, image_path)
    candidate = {
        "id": f"{source_slug}_{slugify(Path(upload.original_name).stem)}",
        "type": "uploaded_image",
        "source_name": upload.original_name,
        "page": None,
        "locator_label": "업로드 이미지",
        "filename": filename,
        "url": f"/api/studio/images/{quote(filename)}",
        "text_preview": Path(upload.original_name).stem,
        "_tokens": tokenize_for_match(Path(upload.original_name).stem),
    }
    candidate["visual_priority"] = visual_candidate_priority(candidate)
    return candidate


def media_asset_to_candidate(asset: dict[str, Any]) -> dict[str, Any]:
    key_findings = asset.get("key_findings") if isinstance(asset.get("key_findings"), list) else []
    text_preview = " ".join(
        [
            str(asset.get("caption") or ""),
            str(asset.get("diagnosis") or ""),
            str(asset.get("modality") or ""),
            " ".join(str(item) for item in key_findings),
            str(asset.get("faculty_note") or ""),
        ]
    ).strip()
    candidate = {
        "id": asset.get("asset_id"),
        "type": "media_bank_asset",
        "source_name": asset.get("original_name") or asset.get("stored_name"),
        "page": None,
        "locator_label": asset.get("modality") or asset.get("asset_type") or "Media Bank",
        "filename": asset.get("stored_name"),
        "url": asset.get("url"),
        "text_preview": text_preview,
        "asset_type": asset.get("asset_type"),
        "modality": asset.get("modality"),
        "diagnosis": asset.get("diagnosis"),
        "approved_for_question_use": asset.get("approved_for_question_use", False),
        "deidentified": asset.get("deidentified", False),
        "_tokens": tokenize_for_match(text_preview),
    }
    candidate["visual_priority"] = max(0.45, visual_candidate_priority(candidate))
    return candidate


def public_visual_candidate(candidate: dict[str, Any], *, match_confidence: float | None = None) -> dict[str, Any]:
    out = {
        "id": candidate.get("id"),
        "type": candidate.get("type"),
        "source_name": candidate.get("source_name"),
        "page": candidate.get("page"),
        "locator_label": candidate.get("locator_label"),
        "visual_priority": candidate.get("visual_priority", visual_candidate_priority(candidate)),
        "picture_count": candidate.get("picture_count", 0),
        "largest_picture_area_ratio": candidate.get("largest_picture_area_ratio", 0),
        "url": candidate.get("url"),
        "text_preview": candidate.get("text_preview", ""),
    }
    if match_confidence is not None:
        out["match_confidence"] = round(match_confidence, 2)
        out["needs_review"] = match_confidence < 0.18
    return out


def score_visual_match(record: dict[str, Any], candidate: dict[str, Any]) -> float:
    pma_solution = record.get("pma_solution") if isinstance(record.get("pma_solution"), dict) else {}
    record_text = " ".join(
        [
            str(record.get("problem") or ""),
            str(record.get("explanation") or ""),
            str(pma_solution.get("reasoning_summary") or ""),
            str(pma_solution.get("source_anchor") or ""),
        ]
    )
    record_tokens = tokenize_for_match(record_text)
    candidate_tokens = candidate.get("_tokens") or set()
    if not record_tokens or not candidate_tokens:
        return 0.0

    overlap = len(record_tokens & candidate_tokens)
    score = overlap / max(6, min(len(record_tokens), len(candidate_tokens)))
    page_num = candidate.get("page")
    if page_num and re.search(rf"(?:page|p\.|페이지|쪽|slide|슬라이드)\s*{page_num}\b", record_text, re.IGNORECASE):
        score += 0.45
    return min(score, 1.0)


def attach_visual_refs(
    records: list[dict[str, Any]],
    candidates: list[dict[str, Any]],
    *,
    force: bool = False,
    max_refs_per_question: int = 1,
) -> list[dict[str, Any]]:
    if not candidates:
        return records
    max_refs = max(1, min(5, int(max_refs_per_question or 1)))

    for idx, record in enumerate(records):
        scored = sorted(
            (
                (
                    score_visual_match(record, candidate),
                    visual_candidate_priority(candidate),
                    candidate,
                )
                for candidate in candidates
            ),
            key=lambda item: (item[0], item[1]),
            reverse=True,
        )
        best_score, best_priority, best_candidate = scored[0]
        if best_score <= 0 and not force:
            continue

        if force:
            ranked = [
                (score if score > 0 else priority or 0.01, priority, candidate)
                for score, priority, candidate in scored
                if score > 0 or priority >= 0.2
            ] or [
                (priority or 0.01, priority, candidate)
                for score, priority, candidate in scored
            ]
            if max_refs == 1 and ranked:
                ranked = [ranked[idx % len(ranked)]]
        else:
            ranked = [(best_score, best_priority, best_candidate)]

        image_refs = []
        seen_ids: set[str] = set()
        for score, priority, candidate in ranked:
            candidate_id = str(candidate.get("id") or candidate.get("url") or "")
            if candidate_id in seen_ids:
                continue
            seen_ids.add(candidate_id)
            confidence = score
            if priority >= 0.45 and confidence < 0.45:
                confidence = max(confidence, priority)
            image_refs.append(public_visual_candidate(candidate, match_confidence=confidence))
            if len(image_refs) >= max_refs:
                break

        if not image_refs:
            continue

        record["image_refs"] = image_refs
        record["image_match_confidence"] = round(max(float(ref.get("match_confidence") or 0) for ref in image_refs), 2)
        if any(ref.get("needs_review") for ref in image_refs):
            record["needs_review"] = True
            reasons = record.get("review_reasons") if isinstance(record.get("review_reasons"), list) else []
            reasons.append("low_image_match_confidence")
            record["review_reasons"] = sorted(set(reasons))
    return records


def filter_visual_candidates(candidates: list[dict[str, Any]], image_policy: str) -> list[dict[str, Any]]:
    if image_policy == "none":
        return []
    if image_policy == "clinical_visuals":
        return [
            candidate
            for candidate in candidates
            if visual_candidate_priority(candidate) >= 0.2
        ]
    return candidates


def normalize_reference_notes(
    records: list[dict[str, Any]],
    *,
    lecture: SavedUpload,
    evidence_uploads: list[SavedUpload],
    reference_policy: str = "uploaded_only",
) -> list[dict[str, Any]]:
    allowed_sources = {lecture.original_name, "강의자료", "lecture"}
    allowed_sources.update(upload.original_name for upload in evidence_uploads)
    allowed_stems = {Path(src).stem for src in allowed_sources}
    allow_unuploaded = reference_policy == "local_open"

    for record in records:
        raw_notes = record.get("reference_notes")
        notes: list[dict[str, Any]] = []
        if isinstance(raw_notes, list):
            for note in raw_notes:
                if not isinstance(note, dict):
                    continue
                source = str(note.get("source") or "").strip()
                if not source:
                    continue
                # Keep source names grounded in uploaded files or explicit lecture wording.
                unuploaded = source not in allowed_sources and source not in allowed_stems
                if unuploaded and not allow_unuploaded:
                    source = lecture.original_name
                basis = str(note.get("basis") or "").strip()
                if unuploaded and allow_unuploaded and "원문 검수 필요" not in basis:
                    basis = f"{basis} (모델 일반지식 기반, 원문 검수 필요)".strip()
                notes.append(
                    {
                        "ref_no": len(notes) + 1,
                        "source": source,
                        "basis": basis,
                        "verification_status": "model_knowledge_unverified" if unuploaded and allow_unuploaded else "uploaded_or_lecture",
                    }
                )

        if not notes:
            refs = record.get("evidence_refs") if isinstance(record.get("evidence_refs"), list) else []
            for ref in refs:
                if not isinstance(ref, dict):
                    continue
                source = str(ref.get("source") or "").strip() or lecture.original_name
                unuploaded = source not in allowed_sources and source not in allowed_stems
                if unuploaded and not allow_unuploaded:
                    source = lecture.original_name
                basis = str(ref.get("basis") or "").strip()
                if unuploaded and allow_unuploaded and "원문 검수 필요" not in basis:
                    basis = f"{basis} (모델 일반지식 기반, 원문 검수 필요)".strip()
                notes.append(
                    {
                        "ref_no": len(notes) + 1,
                        "source": source,
                        "basis": basis,
                        "verification_status": "model_knowledge_unverified" if unuploaded and allow_unuploaded else "uploaded_or_lecture",
                    }
                )

        if not notes:
            notes.append(
                {
                    "ref_no": 1,
                    "source": lecture.original_name,
                    "basis": "업로드 강의자료 기반 문항 초안",
                    "verification_status": "uploaded_or_lecture",
                }
            )

        record["reference_notes"] = notes
        explanation = str(record.get("explanation") or "").strip()
        if "참고문헌:" not in explanation:
            ref_lines = "\n".join(
                f"{note['ref_no']}. {note['source']}: {note.get('basis') or '근거 요약 검수 필요'}"
                for note in notes
            )
            record["explanation"] = f"{explanation}\n\n참고문헌:\n{ref_lines}".strip()
    return records


def generate_anthropic(prompt: str, *, model: str, temperature: float) -> str:
    api_key = os.getenv("ANTHROPIC_API_KEY")
    if not api_key:
        raise RuntimeError("ANTHROPIC_API_KEY가 설정되어 있지 않습니다.")

    response = requests.post(
        "https://api.anthropic.com/v1/messages",
        headers={
            "x-api-key": api_key,
            "anthropic-version": "2023-06-01",
            "content-type": "application/json",
        },
        json={
            "model": model,
            "max_tokens": 8192,
            "temperature": temperature,
            "system": "Return only valid JSON. Do not include markdown fences.",
            "messages": [{"role": "user", "content": prompt}],
        },
        timeout=180,
    )
    if response.status_code >= 400:
        raise RuntimeError(f"Claude API 호출 실패: {response.status_code} {response.text[:500]}")
    payload = response.json()
    return "\n".join(
        block.get("text", "")
        for block in payload.get("content", [])
        if isinstance(block, dict) and block.get("type") == "text"
    ).strip()


def generate_claude_cli(prompt: str, *, model: str, timeout_seconds: int = 420) -> str:
    if not claude_cli_available():
        raise RuntimeError("Claude Code CLI를 찾을 수 없습니다. `claude` 설치/로그인을 먼저 확인해주세요.")
    cmd = [
        "claude",
        "-p",
        "--setting-sources",
        "project,local",
        "--agent",
        "general-purpose",
        "--model",
        model,
        "--output-format",
        "text",
        "--permission-mode",
        "dontAsk",
        "--no-session-persistence",
        "--disallowedTools",
        "Bash,Edit,Read,Write",
        "--append-system-prompt",
        "Return only valid JSON. Do not include markdown fences. Do not read or write files.",
    ]
    result = subprocess.run(
        cmd,
        input=prompt,
        text=True,
        capture_output=True,
        timeout=timeout_seconds,
    )
    if result.returncode != 0:
        raise RuntimeError(f"Claude Code CLI 호출 실패: {result.stderr[:700] or result.stdout[:700]}")
    return result.stdout.strip()


def generate_studio_questions(
    lecture: SavedUpload,
    *,
    subject: str = "General",
    unit: str = "미분류",
    num_questions: int = 5,
    difficulty: str = "보통",
    question_type: str = "clinical_case",
    reference_policy: str = "local_open",
    image_policy: str = "none",
    selected_media_ids: list[str] | None = None,
    image_description: str = "",
    include_tables: bool = False,
    provider: str = "auto",
    style_uploads: list[SavedUpload] | None = None,
    evidence_uploads: list[SavedUpload] | None = None,
    image_uploads: list[SavedUpload] | None = None,
    temperature: float = 0.2,
    model: str | None = None,
) -> dict[str, Any]:
    ensure_studio_dirs()
    style_uploads = style_uploads or []
    evidence_uploads = evidence_uploads or []
    image_uploads = image_uploads or []
    selected_media_ids = selected_media_ids or []
    selected_media_assets = get_media_assets(selected_media_ids)

    lecture_text = extract_saved_upload(lecture)
    if not lecture_text.strip():
        raise ValueError("강의자료에서 텍스트를 추출하지 못했습니다.")

    source_slug = f"{timestamp_slug()}_{slugify(Path(lecture.original_name).stem)}"
    extracted_path = EXTRACTED_DIR / f"{source_slug}.txt"
    prompt_path = GENERATED_DIR / f"{source_slug}.prompt.txt"
    output_path = GENERATED_DIR / f"{source_slug}.questions.json"
    raw_response_path = GENERATED_DIR / f"{source_slug}.model_response.txt"

    extracted_path.write_text(lecture_text, encoding="utf-8")
    raw_visual_candidates: list[dict[str, Any]] = []
    selected_media_candidates = [media_asset_to_candidate(asset) for asset in selected_media_assets]
    raw_visual_candidates.extend(selected_media_candidates)
    if image_policy != "none":
        raw_visual_candidates = extract_visual_candidates(lecture, source_slug, lecture_text)
        raw_visual_candidates = selected_media_candidates + raw_visual_candidates
    if image_policy != "none" and question_type == "image_based":
        for evidence_upload in evidence_uploads:
            try:
                evidence_text_for_images = extract_saved_upload(evidence_upload)
            except Exception:
                evidence_text_for_images = ""
            raw_visual_candidates.extend(
                extract_visual_candidates(
                    evidence_upload,
                    f"{source_slug}_{slugify(Path(evidence_upload.original_name).stem)}",
                    evidence_text_for_images,
                    max_pages=12,
                )
            )
    if image_policy != "none":
        for image_upload in image_uploads:
            image_candidate = direct_image_candidate(image_upload, source_slug)
            if image_candidate:
                raw_visual_candidates.append(image_candidate)
                continue
            try:
                image_text = extract_saved_upload(image_upload)
            except Exception:
                image_text = ""
            raw_visual_candidates.extend(
                extract_visual_candidates(
                    image_upload,
                    f"{source_slug}_{slugify(Path(image_upload.original_name).stem)}",
                    image_text,
                    max_pages=12,
                )
            )
    visual_candidates = filter_visual_candidates(raw_visual_candidates, image_policy)
    public_candidates = [public_visual_candidate(candidate) for candidate in visual_candidates[:12]]

    style_context = build_context_block(
        "style",
        style_uploads,
        max_chars_per_file=5000,
        max_total_chars=10000,
    )
    evidence_context = build_context_block(
        "evidence",
        evidence_uploads,
        max_chars_per_file=8000,
        max_total_chars=20000,
    )

    prompt = build_generation_prompt(
        lecture_text,
        source_name=lecture.original_name,
        subject=subject,
        unit=unit,
        num_questions=num_questions,
        difficulty=difficulty,
        max_chars=30000,
        evidence_context=evidence_context,
    )
    prompt = append_style_context(prompt, style_context)
    prompt = append_question_type_instruction(prompt, question_type, visual_candidates)
    prompt = append_media_instruction(
        prompt,
        visual_candidates,
        selected_media_assets=selected_media_assets,
        image_description=image_description,
    )
    if include_tables:
        prompt = append_table_instruction(prompt)
    prompt = append_reference_instruction(
        prompt,
        lecture,
        evidence_uploads,
        reference_policy=reference_policy,
    )
    prompt_path.write_text(prompt, encoding="utf-8")

    resolved_provider = resolve_studio_provider(provider)
    resolved_model = default_model_for_provider(resolved_provider, model)
    base_response: dict[str, Any] = {
        "source_name": lecture.original_name,
        "provider": resolved_provider,
        "model": resolved_model,
        "subject": subject,
        "unit": unit,
        "num_questions": num_questions,
        "difficulty": difficulty,
        "question_type": question_type,
        "reference_policy": reference_policy,
        "image_policy": image_policy,
        "selected_media_ids": selected_media_ids,
        "include_tables": bool(include_tables),
        "created_at": datetime.now(timezone.utc).isoformat(),
        "pipeline_version": PIPELINE_VERSION,
        "paths": {
            "extracted_text": str(extracted_path),
            "prompt": str(prompt_path),
            "output": str(output_path),
            "raw_response": str(raw_response_path),
        },
        "context": {
            "style_files": [upload.original_name for upload in style_uploads],
            "evidence_files": [upload.original_name for upload in evidence_uploads],
            "image_files": [upload.original_name for upload in image_uploads],
            "selected_media_assets": [
                {
                    "asset_id": asset.get("asset_id"),
                    "modality": asset.get("modality"),
                    "diagnosis": asset.get("diagnosis"),
                    "caption": asset.get("caption"),
                    "url": asset.get("url"),
                }
                for asset in selected_media_assets
            ],
            "image_description": image_description,
            "lecture_chars": len(lecture_text),
            "style_chars": len(style_context),
            "evidence_chars": len(evidence_context),
            "image_candidate_count": len(visual_candidates),
            "raw_image_candidate_count": len(raw_visual_candidates),
        },
        "image_candidates": public_candidates,
    }

    if resolved_provider == "prompt-only":
        packet = {
            **base_response,
            "status": "prompt_ready",
            "question_count": 0,
            "sample": [],
        }
        write_json(output_path, packet)
        return packet

    if resolved_provider == "openai":
        raw_response = generate_openai(
            prompt,
            model=resolved_model,
            temperature=temperature,
        )
    elif resolved_provider == "anthropic":
        raw_response = generate_anthropic(
            prompt,
            model=resolved_model,
            temperature=temperature,
        )
    elif resolved_provider == "claude-cli":
        timeout_seconds = max(360, min(900, 180 + int(num_questions) * 90))
        raw_response = generate_claude_cli(
            prompt,
            model=resolved_model,
            timeout_seconds=timeout_seconds,
        )
    elif resolved_provider == "gemini":
        raw_response = generate_gemini(
            prompt,
            model=resolved_model,
            temperature=temperature,
        )
    else:
        raise ValueError(f"지원하지 않는 provider입니다: {resolved_provider}")

    raw_response_path.write_text(raw_response, encoding="utf-8")
    payload = extract_json_payload(raw_response)
    records = normalize_questions(payload, source_name=lecture.original_name, subject=subject, unit=unit)
    records = normalize_reference_notes(
        records,
        lecture=lecture,
        evidence_uploads=evidence_uploads,
        reference_policy=reference_policy,
    )
    records = attach_visual_refs(
        records,
        visual_candidates,
        force=question_type == "image_based" and bool(visual_candidates),
        max_refs_per_question=3 if question_type == "image_based" else 1,
    )
    write_json(output_path, records)
    archive_path = archive_question_set(
        source_slug,
        {
            **base_response,
            "paths": {
                **base_response["paths"],
                "question_bank": str(QUESTION_BANK_DIR / f"{source_slug}.question_set.json"),
            },
        },
        records,
    )

    sample = [
        {
            "question_id": record.get("question_id"),
            "problem": record.get("problem", "")[:500],
            "options": record.get("options", []),
            "answer": record.get("answer"),
            "explanation": record.get("explanation", "")[:700],
            "needs_review": record.get("needs_review"),
            "review_reasons": record.get("review_reasons", []),
            "evidence_tier": record.get("evidence_tier"),
            "reference_notes": record.get("reference_notes", []),
            "data_table": record.get("data_table"),
            "image_refs": record.get("image_refs", []),
            "image_match_confidence": record.get("image_match_confidence", 0),
        }
        for record in records[:5]
    ]
    base_response["paths"]["question_bank"] = str(archive_path)

    return {
        **base_response,
        "status": "generated",
        "set_id": source_slug,
        "question_count": len(records),
        "sample": sample,
    }
